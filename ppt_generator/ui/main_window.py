"""Main PySide6 window for the technical-proposal PPT desktop workflow."""

from __future__ import annotations

import json
import os
import tempfile
from collections import OrderedDict
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl.utils.cell import get_column_letter
from PySide6.QtCore import QThread, QTimer, Qt, QUrl, Signal
from PySide6.QtGui import QColor, QDesktopServices, QIcon, QPixmap
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QColorDialog,
    QComboBox,
    QDoubleSpinBox,
    QDialog,
    QFileDialog,
    QFormLayout,
    QFrame,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QInputDialog,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMessageBox,
    QPlainTextEdit,
    QProgressBar,
    QPushButton,
    QSplitter,
    QSpinBox,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from .. import __version__
from ..excel_mapper import (
    ExcelMappingRule,
    ExcelWorkbookPreview,
    MappingTarget,
    apply_excel_mappings,
    detect_header_row,
    load_excel_preview,
    preview_rule_value,
    range_shape,
    selection_range,
    suggest_text_mappings,
)
from ..project import AssetRecord, PptProject, SourceRecord, load_project, save_project
from ..preview import (
    OfficePreviewSession,
    ensure_template_thumbnail,
    preview_fingerprint,
    preview_source_slide,
    render_page_preview,
    template_thumbnail_path,
)
from ..module_service import (
    ensure_project_modules,
    rebuild_structure_context,
    slot_specs_for_source_slide,
    sync_legacy_module_state,
)
from ..no_cad_scheme import EquipmentScene
from ..optical_far import apply_optical_far, parse_optical_far
from ..scheme_application import import_no_cad_scene
from ..source_parser import parse_source
from ..template_renderer import TemplateManifest, load_manifest, render_project
from .dialogs import NavigationEditorDialog, TableValueDialog, TextValueDialog
from .auto_solution_editor import AutoSolutionEditor
from .module_editor import ModuleEditor
from .no_cad_scheme_editor import NoCadSchemeEditor
from .scheme_editor import SchemeEditor
from .scheme_visual_lab import SchemeVisualLabWidget
from .slide_preview import SlidePreviewPane
from .styles import APP_QSS


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = PROJECT_ROOT / "templates" / "冲压筒形壳体检测方案NAT6704_v2.pptx"
DEFAULT_MANIFEST = PROJECT_ROOT / "templates" / "NAT6704_v2.template.json"
DEFAULT_DATA = PROJECT_ROOT / "examples" / "NAT6704_v2_test_data.json"
DEFAULT_OUTPUT = PROJECT_ROOT / "output" / "NAT6704_v2_界面生成测试.pptx"
PREVIEW_CACHE_ROOT = PROJECT_ROOT / "output" / ".preview_cache"
FAR_ASSET_ROOT = PROJECT_ROOT / "output" / "far_assets"

ASSET_CATEGORIES = ["未分类", "公司图", "产品图", "设备图", "技术方案图", "检测效果图", "资质图", "其他"]
KIND_NAMES = {"text": "文字", "table": "表格", "image": "图片"}
SLOT_LABELS = {
    "company_name": "公司名称",
    "project_title": "项目标题",
    "project_code": "项目编号",
    "project_date": "方案日期",
    "contact": "联系信息",
    "requirement_title": "需求页标题",
    "inspection_summary": "检测需求概览",
    "special_notes": "特别说明",
    "special_notes_title": "说明标题",
    "product_caption": "产品图片标题",
    "product_image": "产品图片",
    "flow_title": "流程页标题",
    "flow_caption": "流程图标题",
    "flow_result_ok": "OK结果",
    "flow_result_ng": "NG结果",
    "equipment_title": "设备页标题",
    "equipment_description": "设备说明",
    "equipment_caption": "设备图片标题",
    "equipment_image": "设备图片",
    "parameters_title": "参数页标题",
    "equipment_parameters": "设备参数表",
    "parameter_note": "参数备注",
    "parameter_caption": "参数表标题",
    "far_result_title": "检测效果页标题",
    "far_result_camera": "检测效果相机",
    "far_result_lens": "检测效果镜头",
    "far_result_item": "检测效果检测项",
    "far_result_view": "检测效果视角",
    "far_result_image": "检测效果图",
    "far_result_note": "检测效果说明",
    "far_result_caption": "检测效果图题",
    "inspection_items_title": "检测项页标题",
    "inspection_items": "检测项目表",
    "inspection_accuracy_note": "检测精度备注",
    "inspection_items_caption": "检测项表标题",
}


class RenderWorker(QThread):
    succeeded = Signal(str)
    failed = Signal(str)

    def __init__(
        self,
        *,
        project: PptProject,
        overwrite: bool,
        parent=None,
    ):
        super().__init__(parent)
        self.project_data = project.to_dict()
        self.overwrite = overwrite

    def run(self) -> None:
        try:
            project = PptProject.from_dict(self.project_data)
            output = render_project(project, overwrite=self.overwrite)
            self.succeeded.emit(str(output))
        except Exception as exc:
            self.failed.emit(str(exc))


class PreviewWorker(QThread):
    succeeded = Signal(bytes, str, int, str, str, str, bool)
    failed = Signal(str, str, str, str, bool)

    def __init__(
        self,
        *,
        project: PptProject,
        module_id: str,
        slide_id: str,
        cache_key: str,
        office_session: OfficePreviewSession,
        preload: bool = False,
        parent=None,
    ):
        super().__init__(parent)
        self.project_data = project.to_dict()
        self.module_id = module_id
        self.slide_id = slide_id
        self.cache_key = cache_key
        self.office_session = office_session
        self.preload = preload

    def cancel(self) -> None:
        self.requestInterruption()
        self.office_session.cancel_current()

    def run(self) -> None:
        try:
            project = PptProject.from_dict(self.project_data)
            with tempfile.TemporaryDirectory(prefix="kyppt_ui_preview_") as temp_dir:
                output = Path(temp_dir) / "current_page.png"
                path, backend, page_number = render_page_preview(
                    project,
                    self.module_id,
                    self.slide_id,
                    output,
                    office_session=self.office_session,
                    cancelled=self.isInterruptionRequested,
                )
                data = path.read_bytes()
            if not self.isInterruptionRequested():
                self.succeeded.emit(
                    data,
                    backend,
                    page_number,
                    self.module_id,
                    self.slide_id,
                    self.cache_key,
                    self.preload,
                )
        except Exception as exc:
            if not self.isInterruptionRequested():
                self.failed.emit(
                    str(exc),
                    self.module_id,
                    self.slide_id,
                    self.cache_key,
                    self.preload,
                )


class TemplateThumbnailWorker(QThread):
    succeeded = Signal(str, str, int, bool)
    failed = Signal(str, int)

    def __init__(
        self,
        *,
        project: PptProject,
        source_slide: int,
        office_session: OfficePreviewSession,
        parent=None,
    ):
        super().__init__(parent)
        self.project_data = project.to_dict()
        self.source_slide = source_slide
        self.office_session = office_session

    def cancel(self) -> None:
        self.requestInterruption()
        self.office_session.cancel_current()

    def run(self) -> None:
        try:
            project = PptProject.from_dict(self.project_data)
            path, backend, cached = ensure_template_thumbnail(
                project,
                PREVIEW_CACHE_ROOT,
                self.source_slide,
                self.office_session,
                cancelled=self.isInterruptionRequested,
            )
            if not self.isInterruptionRequested():
                self.succeeded.emit(str(path), backend, self.source_slide, cached)
        except Exception as exc:
            if not self.isInterruptionRequested():
                self.failed.emit(str(exc), self.source_slide)


class MainWindow(QMainWindow):
    """DCE-inspired editor with template, module, source, and asset workspaces."""

    def __init__(self):
        super().__init__()
        self.project = PptProject()
        self.manifest: TemplateManifest | None = None
        self.excel_preview: ExcelWorkbookPreview | None = None
        self._project_file: Path | None = None
        self._worker: RenderWorker | None = None
        self._preview_worker: PreviewWorker | None = None
        self._thumbnail_worker: TemplateThumbnailWorker | None = None
        self._thumbnail_failed_slides: set[int] = set()
        self._preview_module_id = ""
        self._preview_slide_id = ""
        self._preview_cache_key = ""
        self._preview_pending = False
        self._preload_queue: list[tuple[str, str, str]] = []
        self._preload_after_current: tuple[str, str] | None = None
        self._office_preview_enabled = (
            os.environ.get("KYPPT_DISABLE_OFFICE_PREVIEW", "") != "1"
            and os.environ.get("QT_QPA_PLATFORM", "").lower()
            not in {"offscreen", "minimal"}
        )
        self._office_preview_session = (
            OfficePreviewSession() if self._office_preview_enabled else None
        )
        self._preview_cache: OrderedDict[
            str, tuple[bytes, str, int]
        ] = OrderedDict()
        self._preview_timer = QTimer(self)
        self._preview_timer.setSingleShot(True)
        self._preview_timer.setInterval(250)
        self._preview_timer.timeout.connect(self._start_page_preview)
        self._field_guard = False
        self._module_guard = False
        self._asset_guard = False
        self._excel_mapping_guard = False

        self.setWindowTitle("KY AI PPT Studio · 技术方案工作台")
        self.setMinimumSize(1280, 720)
        self.resize(1800, 940)
        self.setStyleSheet(APP_QSS)
        self._setup_ui()
        self._load_default_project()

    # ----- UI construction -------------------------------------------------

    def _button(self, text: str, slot, *, primary: bool = False, success: bool = False) -> QPushButton:
        button = QPushButton(text)
        if primary:
            button.setObjectName("primaryButton")
        if success:
            button.setObjectName("successButton")
        button.clicked.connect(slot)
        return button

    def _setup_ui(self) -> None:
        root = QWidget()
        root.setObjectName("rootWidget")
        self.setCentralWidget(root)
        layout = QVBoxLayout(root)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)
        self._setup_topbar(layout)
        self._setup_workspace(layout)

        self.progress = QProgressBar()
        self.progress.setFixedWidth(150)
        self.progress.setRange(0, 0)
        self.progress.hide()
        self.statusBar().addPermanentWidget(self.progress)
        self.statusBar().showMessage("就绪")

    def _setup_topbar(self, parent_layout: QVBoxLayout) -> None:
        topbar = QFrame()
        topbar.setObjectName("topBar")
        topbar.setFixedHeight(60)
        layout = QHBoxLayout(topbar)
        layout.setContentsMargins(18, 8, 18, 8)
        layout.setSpacing(10)

        badge = QLabel("KY")
        badge.setObjectName("brandBadge")
        layout.addWidget(badge)
        brand_text = QVBoxLayout()
        brand_text.setSpacing(0)
        self.brand_title_label = QLabel(f"AI PPT Studio v{__version__}")
        self.brand_title_label.setObjectName("brandTitle")
        subtitle = QLabel("工业设备技术方案 · 本地可编辑输出")
        subtitle.setObjectName("brandSubtitle")
        brand_text.addWidget(self.brand_title_label)
        brand_text.addWidget(subtitle)
        layout.addLayout(brand_text)
        layout.addStretch()

        layout.addWidget(self._button("新建项目", self.new_project))
        layout.addWidget(self._button("打开项目", self.open_project))
        layout.addWidget(self._button("保存项目", self.save_project, primary=True))
        self.open_result_button = self._button("打开结果", self.open_output)
        self.open_result_button.setEnabled(False)
        layout.addWidget(self.open_result_button)
        self.generate_button = self._button("生成 PPT", self.generate_presentation, success=True)
        layout.addWidget(self.generate_button)
        parent_layout.addWidget(topbar)

    def _setup_workspace(self, parent_layout: QVBoxLayout) -> None:
        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.setObjectName("workspaceSplitter")
        splitter.setChildrenCollapsible(False)
        self.workspace_splitter = splitter
        left = QFrame()
        left.setObjectName("leftPanel")
        left_layout = QVBoxLayout(left)
        left_layout.setContentsMargins(14, 14, 8, 14)
        left_layout.setSpacing(10)

        project_group = QGroupBox("当前项目")
        project_form = QFormLayout(project_group)
        self.project_name_edit = QLineEdit()
        self.project_name_edit.setPlaceholderText("例如：NAT6704 筒形壳体检测方案")
        self.project_name_edit.textChanged.connect(self._refresh_summary)
        project_form.addRow("项目名称", self.project_name_edit)
        left_layout.addWidget(project_group)

        self.tabs = QTabWidget()
        self.tabs.setObjectName("workspaceTabs")
        self.tabs.addTab(self._build_content_tab(), "项目内容")
        self.tabs.addTab(self._build_module_tab(), "方案模块")
        self.tabs.addTab(self._build_source_tab(), "Excel / Word")
        self.tabs.addTab(self._build_asset_tab(), "图片素材")
        self.auto_solution_editor = AutoSolutionEditor()
        self.tabs.addTab(self.auto_solution_editor, "自动方案 v2")
        self.scheme_visual_lab_tabs = QTabWidget()
        self.no_cad_scheme_editor = NoCadSchemeEditor()
        self.no_cad_scheme_editor.scheme_committed.connect(
            self._on_no_cad_scheme_committed
        )
        self.no_cad_scheme_editor.workspace_changed.connect(
            self._on_no_cad_workspace_changed
        )
        self.scheme_visual_lab = SchemeVisualLabWidget()
        self.scheme_visual_lab_tabs.addTab(
            self.no_cad_scheme_editor,
            "无CAD逻辑方案",
        )
        self.scheme_visual_lab_tabs.addTab(
            self.scheme_visual_lab,
            "提示词 / SVG实验",
        )
        self.tabs.addTab(self.scheme_visual_lab_tabs, "无CAD方案实验室")
        left_layout.addWidget(self.tabs, 1)
        splitter.addWidget(left)

        right = QFrame()
        right.setObjectName("rightPanel")
        right_layout = QVBoxLayout(right)
        right_layout.setContentsMargins(8, 8, 14, 14)
        right_layout.setSpacing(0)

        review_splitter = QSplitter(Qt.Orientation.Vertical)
        review_splitter.setObjectName("reviewSplitter")
        review_splitter.setChildrenCollapsible(False)
        self.review_splitter = review_splitter

        preview_frame = QFrame()
        preview_frame.setObjectName("previewPanel")
        preview_layout = QVBoxLayout(preview_frame)
        preview_layout.setContentsMargins(0, 0, 0, 0)
        self.slide_preview = SlidePreviewPane()
        self.slide_preview.refresh_requested.connect(
            lambda: self._schedule_page_preview(force=True)
        )
        preview_layout.addWidget(self.slide_preview)
        review_splitter.addWidget(preview_frame)

        review_tabs = QTabWidget()
        review_tabs.setObjectName("reviewTabs")
        self.review_tabs = review_tabs

        output_tab = QWidget()
        output_layout = QGridLayout(output_tab)
        output_layout.setContentsMargins(8, 8, 8, 8)
        output_layout.setHorizontalSpacing(10)
        output_layout.setVerticalSpacing(6)
        output_layout.addWidget(self._build_template_group(), 0, 0, 2, 1)
        output_layout.addWidget(self._build_navigation_style_group(), 0, 1)
        output_layout.addWidget(self._build_metrics_group(), 1, 1)
        output_layout.setColumnStretch(0, 3)
        output_layout.setColumnStretch(1, 2)
        review_tabs.addTab(output_tab, "模板与输出")

        record_tab = QWidget()
        record_layout = QVBoxLayout(record_tab)
        record_layout.setContentsMargins(8, 8, 8, 8)
        record_layout.setSpacing(0)
        record_splitter = QSplitter(Qt.Orientation.Horizontal)
        record_splitter.setChildrenCollapsible(False)
        self.record_splitter = record_splitter

        structure_group = QGroupBox("PPT 结构预览")
        structure_layout = QVBoxLayout(structure_group)
        hint = QLabel("模块、页面增删复制后自动统计；这里显示最终启用的 PPT 结构。")
        hint.setObjectName("mutedLabel")
        structure_layout.addWidget(hint)
        self.structure_list = QListWidget()
        self.structure_list.setObjectName("structureList")
        structure_layout.addWidget(self.structure_list)
        record_splitter.addWidget(structure_group)

        log_group = QGroupBox("运行记录")
        log_layout = QVBoxLayout(log_group)
        self.log_view = QPlainTextEdit()
        self.log_view.setReadOnly(True)
        self.log_view.setMaximumBlockCount(300)
        self.log_view.setPlaceholderText("模板加载、资料解析和 PPT 生成信息会显示在这里。")
        log_layout.addWidget(self.log_view)
        record_splitter.addWidget(log_group)
        record_splitter.setSizes([520, 380])
        record_splitter.setStretchFactor(0, 3)
        record_splitter.setStretchFactor(1, 2)
        record_layout.addWidget(record_splitter)
        review_tabs.addTab(record_tab, "结构与记录")

        review_splitter.addWidget(review_tabs)
        review_splitter.setSizes([560, 270])
        review_splitter.setStretchFactor(0, 3)
        review_splitter.setStretchFactor(1, 1)
        right_layout.addWidget(review_splitter)
        splitter.addWidget(right)

        splitter.setSizes([860, 940])
        splitter.setStretchFactor(0, 5)
        splitter.setStretchFactor(1, 6)
        parent_layout.addWidget(splitter, 1)

    def _build_content_tab(self) -> QWidget:
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(10, 12, 10, 10)
        controls = QHBoxLayout()
        controls.addWidget(self._button("加载数据 JSON", self.load_render_data))
        controls.addWidget(self._button("保存数据 JSON", self.save_render_data))
        controls.addStretch()
        controls.addWidget(self._button("编辑选中字段", self.edit_selected_field, primary=True))
        layout.addLayout(controls)

        self.field_table = QTableWidget(0, 4)
        self.field_table.setObjectName("fieldTable")
        self.field_table.setHorizontalHeaderLabels(["字段", "页码", "类型", "当前值"])
        self.field_table.setAlternatingRowColors(True)
        self.field_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.field_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.field_table.verticalHeader().setVisible(False)
        header = self.field_table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.field_table.itemChanged.connect(self._on_field_item_changed)
        self.field_table.cellDoubleClicked.connect(lambda *_: self.edit_selected_field())
        layout.addWidget(self.field_table)
        return tab

    def _build_module_tab(self) -> QWidget:
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(0, 0, 0, 0)
        self.module_workspace_tabs = QTabWidget()
        self.module_editor = ModuleEditor(SLOT_LABELS)
        self.module_editor.changed.connect(self._on_modules_changed)
        self.module_editor.selection_changed.connect(self._on_preview_selection)
        self.module_editor.message.connect(self._log)
        self.module_workspace_tabs.addTab(self.module_editor, "PPT模块")
        self.scheme_editor = SchemeEditor()
        self.scheme_editor.changed.connect(self._on_scheme_data_changed)
        self.scheme_editor.materialized.connect(self._on_scheme_materialized)
        self.scheme_editor.message.connect(self._log)
        self.module_workspace_tabs.addTab(self.scheme_editor, "设备方案")
        layout.addWidget(self.module_workspace_tabs)
        return tab

    def _build_source_tab(self) -> QWidget:
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(6, 8, 6, 6)
        self.source_inner_tabs = QTabWidget()
        self.source_inner_tabs.addTab(self._build_excel_mapping_panel(), "Excel 映射")
        self.source_inner_tabs.addTab(self._build_generic_source_panel(), "Word / 通用解析")
        layout.addWidget(self.source_inner_tabs)
        return tab

    def _build_excel_mapping_panel(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(8, 10, 8, 8)
        controls = QHBoxLayout()
        controls.addWidget(self._button("选择 Excel", self.import_excel_for_mapping, primary=True))
        self.excel_path_edit = QLineEdit()
        self.excel_path_edit.setReadOnly(True)
        self.excel_path_edit.setPlaceholderText("尚未选择 Excel 文件")
        controls.addWidget(self.excel_path_edit, 1)
        controls.addWidget(QLabel("工作表"))
        self.excel_sheet_combo = QComboBox()
        self.excel_sheet_combo.setMinimumWidth(130)
        self.excel_sheet_combo.currentTextChanged.connect(self._show_excel_sheet)
        controls.addWidget(self.excel_sheet_combo)
        controls.addWidget(QLabel("表头行"))
        self.excel_header_spin = QSpinBox()
        self.excel_header_spin.setRange(1, 1)
        self.excel_header_spin.valueChanged.connect(self._highlight_excel_header)
        controls.addWidget(self.excel_header_spin)
        layout.addLayout(controls)

        self.excel_selection_label = QLabel("先选择 Excel，再框选一个单元格或矩形区域建立映射。")
        self.excel_selection_label.setObjectName("mutedLabel")
        layout.addWidget(self.excel_selection_label)

        far_controls = QHBoxLayout()
        far_hint = QLabel("光学 FAR：自动生成检测效果页，并同步检测项、节拍、相机、镜头和光源")
        far_hint.setObjectName("mutedLabel")
        far_controls.addWidget(far_hint)
        far_controls.addStretch()
        self.far_generate_button = self._button(
            "一键生成检测效果",
            self.generate_inspection_from_far,
            success=True,
        )
        far_controls.addWidget(self.far_generate_button)
        layout.addLayout(far_controls)

        self.excel_preview_table = QTableWidget()
        self.excel_preview_table.setObjectName("excelPreviewTable")
        self.excel_preview_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.excel_preview_table.setSelectionMode(QAbstractItemView.SelectionMode.ContiguousSelection)
        self.excel_preview_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectItems)
        self.excel_preview_table.itemSelectionChanged.connect(self._update_excel_selection_label)
        self.excel_preview_table.verticalHeader().setDefaultSectionSize(30)
        layout.addWidget(self.excel_preview_table, 3)

        mapping_controls = QHBoxLayout()
        mapping_controls.addWidget(self._button("识别表头", self.detect_excel_header))
        mapping_controls.addWidget(self._button("自动建议字段", self.suggest_excel_mappings))
        mapping_controls.addWidget(self._button("添加选区映射", self.add_excel_mapping, primary=True))
        mapping_controls.addWidget(self._button("删除映射", self.remove_excel_mapping))
        mapping_controls.addStretch()
        mapping_controls.addWidget(self._button("保存规则", self.save_excel_mappings))
        mapping_controls.addWidget(self._button("加载规则", self.load_excel_mappings))
        mapping_controls.addWidget(self._button("应用到项目", self.apply_excel_mapping_rules, success=True))
        layout.addLayout(mapping_controls)

        self.excel_mapping_table = QTableWidget(0, 6)
        self.excel_mapping_table.setHorizontalHeaderLabels(
            ["启用", "工作表", "来源范围", "PPT目标", "写入方式", "数据预览"]
        )
        self.excel_mapping_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.excel_mapping_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.excel_mapping_table.verticalHeader().setVisible(False)
        mapping_header = self.excel_mapping_table.horizontalHeader()
        for column in range(5):
            mapping_header.setSectionResizeMode(column, QHeaderView.ResizeMode.ResizeToContents)
        mapping_header.setSectionResizeMode(5, QHeaderView.ResizeMode.Stretch)
        self.excel_mapping_table.itemChanged.connect(self._on_excel_mapping_item_changed)
        layout.addWidget(self.excel_mapping_table, 2)
        return panel

    def _build_generic_source_panel(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(8, 10, 8, 8)
        controls = QHBoxLayout()
        controls.addWidget(self._button("导入 Excel / Word", self.import_sources, primary=True))
        controls.addWidget(self._button("解析选中", self.parse_selected_source))
        controls.addWidget(self._button("移除", self.remove_selected_source))
        layout.addLayout(controls)
        self.source_list = QListWidget()
        self.source_list.setMaximumHeight(150)
        self.source_list.currentRowChanged.connect(self._show_source_preview)
        layout.addWidget(self.source_list)
        self.source_preview = QPlainTextEdit()
        self.source_preview.setPlaceholderText("选择资料并点击“解析选中”，提取的文字和表格会显示在这里。")
        layout.addWidget(self.source_preview, 1)
        apply_row = QHBoxLayout()
        apply_row.addWidget(QLabel("应用到文字字段"))
        self.source_target_combo = QComboBox()
        apply_row.addWidget(self.source_target_combo, 1)
        apply_row.addWidget(self._button("填入选中文字", self.apply_source_to_slot, primary=True))
        layout.addLayout(apply_row)
        return panel

    def _build_asset_tab(self) -> QWidget:
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(10, 12, 10, 10)
        controls = QHBoxLayout()
        controls.addWidget(self._button("导入图片", self.import_images, primary=True))
        controls.addWidget(self._button("移除选中", self.remove_selected_asset))
        controls.addStretch()
        layout.addLayout(controls)
        self.asset_table = QTableWidget(0, 4)
        self.asset_table.setHorizontalHeaderLabels(["预览", "文件", "分类", "PPT用途"])
        self.asset_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.asset_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.asset_table.verticalHeader().setVisible(False)
        header = self.asset_table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.ResizeToContents)
        layout.addWidget(self.asset_table)
        hint = QLabel("分类用于素材管理；分配到“产品图片/设备图片”后会在生成时替换对应模板图片。")
        hint.setObjectName("mutedLabel")
        layout.addWidget(hint)
        return tab

    def _build_template_group(self) -> QGroupBox:
        group = QGroupBox("模板与输出")
        layout = QGridLayout(group)
        self.template_path_edit = QLineEdit()
        self.template_path_edit.setReadOnly(True)
        self.manifest_path_edit = QLineEdit()
        self.manifest_path_edit.setReadOnly(True)
        self.output_path_edit = QLineEdit()
        self.output_path_edit.setReadOnly(True)
        layout.addWidget(QLabel("PPT模板"), 0, 0)
        layout.addWidget(self.template_path_edit, 0, 1)
        layout.addWidget(self._button("选择", self.choose_template), 0, 2)
        layout.addWidget(QLabel("模板配置"), 1, 0)
        layout.addWidget(self.manifest_path_edit, 1, 1)
        layout.addWidget(self._button("选择", self.choose_manifest), 1, 2)
        layout.addWidget(QLabel("输出文件"), 2, 0)
        layout.addWidget(self.output_path_edit, 2, 1)
        layout.addWidget(self._button("选择", self.choose_output), 2, 2)
        layout.setColumnStretch(1, 1)
        return group

    def _build_navigation_style_group(self) -> QGroupBox:
        group = QGroupBox("导航样式")
        layout = QGridLayout(group)

        self.navigation_height_spin = QDoubleSpinBox()
        self.navigation_height_spin.setRange(0.42, 0.72)
        self.navigation_height_spin.setDecimals(2)
        self.navigation_height_spin.setSingleStep(0.02)
        self.navigation_height_spin.setSuffix(" in")
        self.navigation_height_spin.valueChanged.connect(
            self._on_navigation_height_changed
        )
        layout.addWidget(QLabel("导航高度"), 0, 0)
        layout.addWidget(self.navigation_height_spin, 0, 1)

        self.navigation_font_size_combo = QComboBox()
        self.navigation_font_size_combo.addItem("自动（随高度）", None)
        for font_size in range(9, 17):
            self.navigation_font_size_combo.addItem(f"{font_size} pt", float(font_size))
        self.navigation_font_size_combo.currentIndexChanged.connect(
            self._on_navigation_font_size_changed
        )
        layout.addWidget(QLabel("导航字体"), 1, 0)
        layout.addWidget(self.navigation_font_size_combo, 1, 1)

        self.navigation_color_button = self._button(
            "选择 #FFFFFF",
            self.choose_navigation_background,
        )
        self.navigation_background_label = QLabel("当前栏目背景")
        layout.addWidget(self.navigation_background_label, 2, 0)
        layout.addWidget(self.navigation_color_button, 2, 1)

        layout.addWidget(QLabel("导航栏目"), 3, 0)
        navigation_row = QHBoxLayout()
        self.navigation_items_label = QLabel()
        self.navigation_items_label.setObjectName("mutedLabel")
        self.navigation_items_label.setWordWrap(True)
        navigation_row.addWidget(self.navigation_items_label, 1)
        self.edit_navigation_button = self._button(
            "编辑栏目",
            self.edit_navigation_items,
        )
        navigation_row.addWidget(self.edit_navigation_button)
        layout.addLayout(navigation_row, 3, 1)

        hint = QLabel("支持1～7项并分配PPT模块；自动字号随导航高度变化。")
        hint.setObjectName("mutedLabel")
        hint.setWordWrap(True)
        layout.addWidget(hint, 4, 0, 1, 2)
        layout.setColumnStretch(1, 1)
        return group

    def _build_metrics_group(self) -> QGroupBox:
        group = QGroupBox("生成检查")
        layout = QHBoxLayout(group)
        self.metric_slides = self._metric(layout, "0", "预计页数")
        self.metric_modules = self._metric(layout, "0", "启用模块")
        self.metric_fields = self._metric(layout, "0", "已填字段")
        self.metric_sources = self._metric(layout, "0", "导入资料")
        self.metric_assets = self._metric(layout, "0", "图片素材")
        return group

    def _metric(self, layout: QHBoxLayout, value: str, name: str) -> QLabel:
        container = QVBoxLayout()
        value_label = QLabel(value)
        value_label.setObjectName("metricValue")
        value_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        name_label = QLabel(name)
        name_label.setObjectName("metricName")
        name_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        container.addWidget(value_label)
        container.addWidget(name_label)
        layout.addLayout(container, 1)
        return value_label

    # ----- project and template state -------------------------------------

    def _update_navigation_color_button(self, color_value: str) -> None:
        color = QColor(color_value)
        foreground = "#FFFFFF" if color.lightness() < 128 else "#1F2933"
        self.navigation_color_button.setText(f"选择 {color_value.upper()}")
        self.navigation_color_button.setStyleSheet(
            "QPushButton {"
            f"background-color: {color_value}; color: {foreground};"
            "border: 1px solid #9AA6AF;"
            "}"
        )

    def _sync_navigation_style_ui(self) -> None:
        style = self.project.presentation_style
        self.navigation_height_spin.blockSignals(True)
        self.navigation_height_spin.setValue(style.navigation_height)
        self.navigation_height_spin.blockSignals(False)
        self.navigation_font_size_combo.blockSignals(True)
        target_index = self.navigation_font_size_combo.findData(
            style.navigation_font_size
        )
        self.navigation_font_size_combo.setCurrentIndex(max(0, target_index))
        self.navigation_font_size_combo.blockSignals(False)
        self._update_navigation_color_button(style.navigation_background)
        self._update_navigation_items_summary()

    def _update_navigation_items_summary(self) -> None:
        items = self.project.presentation_style.navigation_items
        self.navigation_items_label.setText(" ｜ ".join(item.name for item in items))
        module_names = {
            str(module.get("key") or ""): str(module.get("name") or module.get("key") or "")
            for module in (self.manifest.modules if self.manifest else [])
        }
        details = []
        for item in items:
            assigned = "、".join(
                module_names.get(key, key) for key in item.module_keys
            ) or "未分配模块"
            details.append(f"{item.name}：{assigned}")
        self.navigation_items_label.setToolTip("\n".join(details))

    def _refresh_after_navigation_style_change(self) -> None:
        self._apply_navigation_preview_overlay()
        self._preview_cache.clear()
        self._preview_cache_key = ""
        self._preload_queue.clear()
        self._preload_after_current = None
        self._schedule_page_preview(force=True)

    def _on_navigation_height_changed(self, value: float) -> None:
        self.project.presentation_style.navigation_height = round(float(value), 2)
        self._refresh_after_navigation_style_change()

    def _on_navigation_font_size_changed(self, _index: int) -> None:
        value = self.navigation_font_size_combo.currentData()
        self.project.presentation_style.navigation_font_size = (
            None if value is None else float(value)
        )
        self._refresh_after_navigation_style_change()

    def choose_navigation_background(self) -> None:
        selected = QColorDialog.getColor(
            QColor(self.project.presentation_style.navigation_background),
            self,
            "选择当前栏目背景颜色",
        )
        if not selected.isValid():
            return
        color_value = selected.name().upper()
        self.project.presentation_style.navigation_background = color_value
        self._update_navigation_color_button(color_value)
        self._refresh_after_navigation_style_change()

    def _apply_navigation_preview_overlay(self) -> bool:
        module = next(
            (
                item
                for item in self.project.modules
                if item.id == self._preview_module_id
            ),
            None,
        )
        if module is None:
            return False
        style = self.project.presentation_style
        return self.slide_preview.apply_navigation_overlay(
            [item.name for item in style.navigation_items],
            style.navigation_index_for(module.template_module_key),
            style.navigation_height,
            style.navigation_background,
            style.resolved_navigation_font_size(),
        )

    def edit_navigation_items(self) -> None:
        modules = [
            (
                str(module.get("key") or ""),
                str(module.get("name") or module.get("key") or ""),
            )
            for module in (self.manifest.modules if self.manifest else [])
            if module.get("key")
        ]
        dialog = NavigationEditorDialog(
            self.project.presentation_style.navigation_items,
            modules,
            self,
        )
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        self.project.presentation_style.navigation_items = dialog.navigation_items
        self.project.presentation_style.validate()
        self._update_navigation_items_summary()
        self._log("顶部导航栏目已更新")
        self._refresh_after_navigation_style_change()

    def _load_default_project(self) -> None:
        values: dict[str, Any] = {}
        if DEFAULT_DATA.is_file():
            values = json.loads(DEFAULT_DATA.read_text(encoding="utf-8"))
        self.project = PptProject(
            project_name="NAT6704 筒形壳体检测技术方案",
            template_path=str(DEFAULT_TEMPLATE),
            manifest_path=str(DEFAULT_MANIFEST),
            output_path=str(DEFAULT_OUTPUT),
            values=values,
        )
        self._project_file = None
        self._load_manifest_into_ui(reset_modules=True)
        self._log("已加载 NAT6704 默认模板和 M2 测试数据。")

    def _load_manifest_into_ui(self, *, reset_modules: bool) -> None:
        self._preview_cache.clear()
        self._thumbnail_failed_slides.clear()
        self._preload_queue.clear()
        self._preload_after_current = None
        self.manifest = load_manifest(self.project.manifest_path)
        ensure_project_modules(
            self.project,
            self.manifest,
            reset=reset_modules,
        )
        known_slots = {slot["key"] for slot in self.manifest.slots}
        self.project.values = {
            key: value for key, value in self.project.values.items() if key in known_slots
        }
        self.project.excel_mappings = [
            rule for rule in self.project.excel_mappings if rule.target_slot in known_slots
        ]
        for asset in self.project.assets:
            if asset.slot_key not in known_slots:
                asset.slot_key = ""

        self.project_name_edit.setText(self.project.project_name)
        self.template_path_edit.setText(self.project.template_path)
        self.manifest_path_edit.setText(self.project.manifest_path)
        self.output_path_edit.setText(self.project.output_path)
        self._sync_navigation_style_ui()
        self.module_editor.set_state(self.project, self.manifest)
        self.scheme_editor.set_state(self.project, self.manifest)
        self.no_cad_scheme_editor.set_project_context(
            self.project.project_id,
            self.project.no_cad_scene,
            self.project.ai_image_batches,
        )
        self._refresh_structure()
        self._refresh_field_table()
        self._refresh_source_targets()
        self._refresh_sources()
        self._restore_excel_state()
        self._refresh_asset_table()
        self._refresh_summary()

    def _on_template_thumbnail_ready(
        self,
        _path: str,
        backend: str,
        source_slide: int,
        cached: bool,
    ) -> None:
        if not cached:
            self._log(f"模板第 {source_slide} 页缩略图已缓存 · {backend}")
        try:
            current_source = preview_source_slide(
                self.project,
                self._preview_module_id,
                self._preview_slide_id,
            )
        except Exception:
            return
        if current_source == source_slide:
            self._show_template_thumbnail_for_current()

    def _on_template_thumbnail_failed(
        self,
        message: str,
        source_slide: int,
    ) -> None:
        self._thumbnail_failed_slides.add(source_slide)
        self._log(f"模板第 {source_slide} 页缩略图生成失败：{message}")

    def _on_template_thumbnail_worker_finished(self) -> None:
        worker = self._thumbnail_worker
        if worker:
            worker.deleteLater()
        self._thumbnail_worker = None
        if self._preview_module_id and self._preview_slide_id:
            self._preview_pending = False
            self._preview_timer.start(0)

    def _ensure_current_template_thumbnail(self) -> bool:
        """Ensure the first-stage image exists; return True while waiting."""
        if self._show_template_thumbnail_for_current():
            return False
        try:
            source_slide = preview_source_slide(
                self.project,
                self._preview_module_id,
                self._preview_slide_id,
            )
        except Exception:
            return False
        if source_slide in self._thumbnail_failed_slides:
            return False
        if self._thumbnail_worker and self._thumbnail_worker.isRunning():
            self._preview_pending = True
            return True
        if not self._office_preview_enabled or not self._office_preview_session:
            return False
        worker = TemplateThumbnailWorker(
            project=self.project,
            source_slide=source_slide,
            office_session=self._office_preview_session,
            parent=self,
        )
        worker.succeeded.connect(self._on_template_thumbnail_ready)
        worker.failed.connect(self._on_template_thumbnail_failed)
        worker.finished.connect(self._on_template_thumbnail_worker_finished)
        self._thumbnail_worker = worker
        self._preview_pending = True
        self.slide_preview.set_loading(self._current_preview_name())
        worker.start()
        return True

    def _current_preview_name(self) -> str:
        module = next(
            (
                item
                for item in self.project.modules
                if item.id == self._preview_module_id
            ),
            None,
        )
        slide = next(
            (
                item
                for item in module.slides
                if item.id == self._preview_slide_id
            ),
            None,
        ) if module else None
        return slide.title if slide and slide.title else "当前页面"

    def _show_template_thumbnail_for_current(self) -> bool:
        if not self._preview_module_id or not self._preview_slide_id:
            return False
        if self._preview_cache_key in self._preview_cache:
            return False
        try:
            source_slide = preview_source_slide(
                self.project,
                self._preview_module_id,
                self._preview_slide_id,
            )
            path = template_thumbnail_path(
                self.project,
                PREVIEW_CACHE_ROOT,
                source_slide,
            )
            data = path.read_bytes()
            if not data.startswith(b"\x89PNG\r\n\x1a\n"):
                return False
        except (OSError, RuntimeError, ValueError):
            return False
        self.slide_preview.set_base_preview(
            data,
            source_slide,
            self._current_preview_name(),
        )
        self._apply_navigation_preview_overlay()
        return True

    def _slot_by_key(self, key: str) -> dict[str, Any] | None:
        if not self.manifest:
            return None
        return next((slot for slot in self.manifest.slots if slot["key"] == key), None)

    def _slot_label(self, slot: dict[str, Any]) -> str:
        key = slot["key"]
        if key.startswith("flow_step_"):
            return f"流程步骤 {key.rsplit('_', 1)[-1]}"
        return str(slot.get("label") or SLOT_LABELS.get(key) or key)

    def _module_by_key(self) -> dict[str, dict[str, Any]]:
        if not self.manifest:
            return {}
        return {module["key"]: module for module in self.manifest.modules}

    def _collect_ui_state(self) -> None:
        self.project.project_name = self.project_name_edit.text().strip() or "未命名技术方案"
        self.project.template_path = self.template_path_edit.text()
        self.project.manifest_path = self.manifest_path_edit.text()
        self.project.output_path = self.output_path_edit.text()
        self.project.presentation_style.navigation_height = round(
            self.navigation_height_spin.value(), 2
        )
        self.project.presentation_style.validate()
        no_cad_snapshot = self.no_cad_scheme_editor.workspace_snapshot()
        if no_cad_snapshot.get("projectId") == self.project.project_id:
            self.project.no_cad_scene = deepcopy(no_cad_snapshot["scene"])
            self.project.ai_image_batches = deepcopy(
                no_cad_snapshot["aiImageBatches"]
            )
        sync_legacy_module_state(self.project)

    def new_project(self) -> None:
        answer = QMessageBox.question(
            self,
            "新建项目",
            "将重新加载默认模板和测试数据。未保存的修改会丢失，是否继续？",
        )
        if answer == QMessageBox.StandardButton.Yes:
            self._load_default_project()

    def open_project(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "打开 KY PPT 项目",
            str(PROJECT_ROOT),
            "KY PPT 项目 (*.kyppt.json *.json);;所有文件 (*)",
        )
        if not path:
            return
        try:
            project = load_project(path)
            if not Path(project.template_path).is_file() or not Path(project.manifest_path).is_file():
                raise FileNotFoundError("项目引用的模板或模板配置不存在")
            self.project = project
            self._project_file = Path(path)
            self._load_manifest_into_ui(reset_modules=False)
            self._log(f"已打开项目：{path}")
            self.statusBar().showMessage("项目已打开", 4000)
        except Exception as exc:
            self._error("打开项目失败", str(exc))

    def save_project(self) -> None:
        self._collect_ui_state()
        path = str(self._project_file) if self._project_file else ""
        if not path:
            path, _ = QFileDialog.getSaveFileName(
                self,
                "保存 KY PPT 项目",
                str(PROJECT_ROOT / "projects" / f"{self.project.project_name}.kyppt.json"),
                "KY PPT 项目 (*.kyppt.json);;JSON (*.json)",
            )
        if not path:
            return
        try:
            self._project_file = save_project(self.project, path)
            self._log(f"项目已保存：{self._project_file}")
            self.statusBar().showMessage("项目已保存", 4000)
        except Exception as exc:
            self._error("保存项目失败", str(exc))

    def choose_template(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择 PPTX 模板",
            str(PROJECT_ROOT / "templates"),
            "PowerPoint 模板 (*.pptx)",
        )
        if not path:
            return
        manifest_path = self._find_manifest(Path(path))
        if manifest_path is None:
            selected, _ = QFileDialog.getOpenFileName(
                self,
                "选择与模板对应的配置 JSON",
                str(Path(path).parent),
                "模板配置 (*.template.json *.json)",
            )
            if not selected:
                return
            manifest_path = Path(selected)
        try:
            self.project.template_path = path
            self.project.manifest_path = str(manifest_path)
            self._load_manifest_into_ui(reset_modules=True)
            self._log(f"已选择模板：{path}")
        except Exception as exc:
            self._error("模板加载失败", str(exc))

    def _find_manifest(self, template_path: Path) -> Path | None:
        for candidate in template_path.parent.glob("*.template.json"):
            try:
                if load_manifest(candidate).template_filename == template_path.name:
                    return candidate
            except Exception:
                continue
        return None

    def choose_manifest(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择模板配置",
            str(Path(self.project.manifest_path).parent if self.project.manifest_path else PROJECT_ROOT),
            "模板配置 (*.template.json *.json)",
        )
        if not path:
            return
        try:
            self.project.manifest_path = path
            self._load_manifest_into_ui(reset_modules=True)
            self._log(f"已加载模板配置：{path}")
        except Exception as exc:
            self._error("模板配置加载失败", str(exc))

    def choose_output(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self,
            "选择 PPT 输出文件",
            self.project.output_path or str(DEFAULT_OUTPUT),
            "PowerPoint 文件 (*.pptx)",
        )
        if path:
            self.project.output_path = path
            self.output_path_edit.setText(path)

    # ----- field editing ---------------------------------------------------

    def _value_summary(self, slot: dict[str, Any]) -> str:
        value = self.project.values.get(slot["key"])
        if slot["kind"] == "table":
            if isinstance(value, list) and value:
                columns = len(value[0]) if isinstance(value[0], list) else 0
                return f"{len(value)} × {columns} 表格（双击编辑）"
            return "未填写（双击编辑）"
        if slot["kind"] == "image":
            return Path(value).name if isinstance(value, str) and value else "未选择图片（双击选择）"
        return str(value or "").replace("\n", "  ↵  ")

    def _refresh_field_table(self) -> None:
        self._field_guard = True
        try:
            slots = list(self.manifest.slots) if self.manifest else []
            self.field_table.setRowCount(len(slots))
            for row, slot in enumerate(slots):
                label_item = QTableWidgetItem(self._slot_label(slot))
                label_item.setData(Qt.ItemDataRole.UserRole, slot["key"])
                label_item.setFlags(label_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                page_item = QTableWidgetItem(str(slot["slide"]))
                page_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                page_item.setFlags(page_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                kind_item = QTableWidgetItem(KIND_NAMES[slot["kind"]])
                kind_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                kind_item.setFlags(kind_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                value_item = QTableWidgetItem(self._value_summary(slot))
                value_item.setToolTip(str(self.project.values.get(slot["key"], "")))
                if slot["kind"] != "text":
                    value_item.setFlags(value_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.field_table.setItem(row, 0, label_item)
                self.field_table.setItem(row, 1, page_item)
                self.field_table.setItem(row, 2, kind_item)
                self.field_table.setItem(row, 3, value_item)
                self.field_table.setRowHeight(row, 38)
        finally:
            self._field_guard = False
        if hasattr(self, "slide_preview"):
            self._schedule_page_preview()

    def _on_field_item_changed(self, item: QTableWidgetItem) -> None:
        if self._field_guard or item.column() != 3:
            return
        key_item = self.field_table.item(item.row(), 0)
        key = key_item.data(Qt.ItemDataRole.UserRole) if key_item else ""
        slot = self._slot_by_key(str(key))
        if slot and slot["kind"] == "text":
            self.project.values[str(key)] = item.text().replace("  ↵  ", "\n")
            self._refresh_summary()
            self._schedule_page_preview()

    def edit_selected_field(self) -> None:
        row = self.field_table.currentRow()
        if row < 0:
            QMessageBox.information(self, "编辑字段", "请先选择一个字段。")
            return
        key = self.field_table.item(row, 0).data(Qt.ItemDataRole.UserRole)
        slot = self._slot_by_key(str(key))
        if not slot:
            return
        value = self.project.values.get(str(key), "")
        title = f"编辑 · {self._slot_label(slot)}"
        if slot["kind"] == "text":
            dialog = TextValueDialog(title, str(value or ""), slot.get("max_chars"), self)
            if dialog.exec():
                self.project.values[str(key)] = dialog.value
        elif slot["kind"] == "table":
            rows = int(slot.get("rows") or (len(value) if isinstance(value, list) else 0))
            columns = int(
                slot.get("columns")
                or (
                    len(value[0])
                    if isinstance(value, list) and value and isinstance(value[0], list)
                    else 0
                )
            )
            if not rows or not columns:
                self._error("无法编辑表格", "模板配置没有记录表格行列数。")
                return
            dialog = TableValueDialog(title, value, rows, columns, self)
            if dialog.exec():
                self.project.values[str(key)] = dialog.value
        else:
            path, _ = QFileDialog.getOpenFileName(
                self,
                "选择替换图片",
                str(PROJECT_ROOT / "assets"),
                "图片 (*.png *.jpg *.jpeg *.bmp *.webp)",
            )
            if path:
                self.project.values[str(key)] = path
        self._refresh_field_table()
        self._refresh_summary()

    def load_render_data(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, "加载渲染数据", str(PROJECT_ROOT / "examples"), "JSON (*.json)"
        )
        if not path:
            return
        try:
            raw = json.loads(Path(path).read_text(encoding="utf-8"))
            if not isinstance(raw, dict):
                raise ValueError("数据 JSON 顶层必须是对象")
            known = {slot["key"] for slot in self.manifest.slots} if self.manifest else set()
            image_keys = {
                slot["key"] for slot in self.manifest.slots if slot["kind"] == "image"
            } if self.manifest else set()
            for key in image_keys & set(raw):
                value = raw[key]
                if isinstance(value, str) and value and not Path(value).is_absolute():
                    raw[key] = str((Path(path).parent / value).resolve())
            unknown = sorted(set(raw) - known)
            self.project.values.update({key: value for key, value in raw.items() if key in known})
            self._refresh_field_table()
            self._refresh_summary()
            self._log(f"已加载数据：{path}")
            if unknown:
                QMessageBox.warning(self, "部分字段未载入", f"模板未配置：{', '.join(unknown)}")
        except Exception as exc:
            self._error("加载数据失败", str(exc))

    def save_render_data(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self, "保存渲染数据", str(PROJECT_ROOT / "examples" / "project_data.json"), "JSON (*.json)"
        )
        if not path:
            return
        try:
            Path(path).write_text(
                json.dumps(self.project.values, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            self._log(f"渲染数据已保存：{path}")
        except Exception as exc:
            self._error("保存数据失败", str(exc))

    # ----- modules ---------------------------------------------------------

    def _refresh_module_list(self) -> None:
        if self.manifest:
            self.module_editor.set_state(self.project, self.manifest)
            self.scheme_editor.set_state(self.project, self.manifest)

    def _on_scheme_data_changed(self) -> None:
        self._refresh_asset_table()
        self._refresh_summary()

    def _on_scheme_materialized(self) -> None:
        selection_anchor = self.module_editor.selection_anchor()
        if self.manifest:
            self.module_editor.set_state(
                self.project,
                self.manifest,
                selection_anchor=selection_anchor,
            )
        self._on_modules_changed()

    def _on_no_cad_scheme_committed(self, scene_data: object) -> None:
        if not isinstance(scene_data, dict):
            self._error("同步无CAD方案失败", "Scene 数据格式无效。")
            return
        try:
            scene = EquipmentScene.from_dict(scene_data)
            result = import_no_cad_scene(self.project, scene)
        except Exception as exc:
            self._error("同步无CAD方案失败", str(exc))
            return
        if self.manifest:
            self.scheme_editor.set_state(self.project, self.manifest)
        self._refresh_asset_table()
        self._refresh_summary()
        self.tabs.setCurrentIndex(1)
        self.module_workspace_tabs.setCurrentWidget(self.scheme_editor)
        message = (
            f"无CAD结构已同步到正式设备方案：{result.equipment_modules} 个模块，"
            f"{result.image_targets} 个图片目标，尚缺 {result.pending_images} 张采用图。"
        )
        self._log(message)

    def _on_no_cad_workspace_changed(self, payload: object) -> None:
        if not isinstance(payload, dict):
            return
        project_id = str(payload.get("projectId") or "")
        if not hasattr(self, "project") or project_id != self.project.project_id:
            return
        scene_data = payload.get("scene")
        batch_records = payload.get("aiImageBatches")
        if not isinstance(scene_data, dict) or not isinstance(batch_records, list):
            return
        self.project.no_cad_scene = deepcopy(scene_data)
        self.project.ai_image_batches = deepcopy(batch_records)
        if not bool(payload.get("candidateHistoryChanged")):
            return
        if self._project_file is None:
            self._log("候选图已绑定当前新项目；请保存项目以持久化批次索引。")
            self.statusBar().showMessage("候选图已绑定，请保存项目", 6000)
            return
        try:
            self._collect_ui_state()
            save_project(self.project, self._project_file)
            self._log(f"候选图批次已自动写入项目：{self._project_file}")
            self.statusBar().showMessage("候选图批次已自动保存", 5000)
        except Exception as exc:
            self._error("候选图批次保存失败", str(exc))

    def _capture_module_state(self, *_args) -> None:
        sync_legacy_module_state(self.project)

    def _set_all_modules(self, enabled: bool) -> None:
        for module in self.project.modules:
            module.enabled = enabled
        sync_legacy_module_state(self.project)
        self._refresh_module_list()
        self._on_modules_changed()

    def _move_module(self, offset: int) -> None:
        self.module_editor.move_selected(offset)

    def _on_modules_changed(self) -> None:
        sync_legacy_module_state(self.project)
        self._preload_queue.clear()
        self._preload_after_current = None
        self._refresh_structure()
        self._refresh_summary()
        self._schedule_page_preview()

    def _on_preview_selection(self, module_id: str, slide_id: str) -> None:
        self._preview_timer.stop()
        self._preview_module_id = module_id
        self._preview_slide_id = slide_id
        self._preview_cache_key = ""
        if not module_id or not slide_id:
            self.slide_preview.clear_selection()
            return

        module = next(
            (item for item in self.project.modules if item.id == module_id),
            None,
        )
        slide = next(
            (item for item in module.slides if item.id == slide_id),
            None,
        ) if module else None
        if module is None or slide is None:
            self.slide_preview.clear_selection("当前页面已不存在")
            return

        self.slide_preview.set_selection(
            module.id,
            slide.id,
            module.name,
            slide.title or "未命名页面",
        )
        if not module.enabled:
            self.slide_preview.set_error(
                "当前模块未启用，不属于最终PPT；启用模块后可生成预览。"
            )
            return
        self._show_template_thumbnail_for_current()
        self._schedule_page_preview()

    def _schedule_page_preview(self, *, force: bool = False) -> None:
        if not self._preview_module_id or not self._preview_slide_id:
            return
        module = next(
            (
                item
                for item in self.project.modules
                if item.id == self._preview_module_id
            ),
            None,
        )
        if module is None or not module.enabled:
            self.slide_preview.set_error(
                "当前模块未启用，不属于最终PPT；启用模块后可生成预览。"
            )
            return
        try:
            cache_key = preview_fingerprint(
                self.project,
                self._preview_module_id,
                self._preview_slide_id,
            )
        except Exception as exc:
            self.slide_preview.set_error(str(exc))
            return
        self._preview_cache_key = cache_key
        if not force and cache_key in self._preview_cache:
            data, backend, page_number = self._preview_cache.pop(cache_key)
            self._preview_cache[cache_key] = (data, backend, page_number)
            self._preview_timer.stop()
            self.slide_preview.set_preview(
                data,
                backend,
                page_number,
                cached=True,
            )
            return
        self.slide_preview.set_loading(self._current_preview_name())
        if not self._office_preview_enabled:
            self.slide_preview.set_error(
                "当前为自动化测试显示模式，已暂停调用 PowerPoint/WPS。"
            )
            return
        self._preview_timer.start()

    def _start_page_preview(self) -> None:
        if not self._preview_module_id or not self._preview_slide_id:
            return
        if self._preview_worker and self._preview_worker.isRunning():
            self._preview_pending = True
            return
        if not self._office_preview_session or not self._preview_cache_key:
            return
        if self._ensure_current_template_thumbnail():
            return

        self._preview_pending = False
        self.slide_preview.set_loading(self._current_preview_name())
        worker = PreviewWorker(
            project=self.project,
            module_id=self._preview_module_id,
            slide_id=self._preview_slide_id,
            cache_key=self._preview_cache_key,
            office_session=self._office_preview_session,
            preload=False,
            parent=self,
        )
        worker.succeeded.connect(self._on_preview_succeeded)
        worker.failed.connect(self._on_preview_failed)
        worker.finished.connect(self._on_preview_finished)
        self._preview_worker = worker
        worker.start()

    def _on_preview_succeeded(
        self,
        data: bytes,
        backend: str,
        page_number: int,
        module_id: str,
        slide_id: str,
        cache_key: str,
        preload: bool,
    ) -> None:
        self._preview_cache[cache_key] = (data, backend, page_number)
        self._preview_cache.move_to_end(cache_key)
        while len(self._preview_cache) > 32:
            self._preview_cache.popitem(last=False)
        if (
            module_id == self._preview_module_id
            and slide_id == self._preview_slide_id
            and cache_key == self._preview_cache_key
        ):
            self.slide_preview.set_preview(data, backend, page_number)
            if preload:
                self._preview_pending = False
                self._preview_timer.stop()
            else:
                self._preload_after_current = (module_id, slide_id)

    def _on_preview_failed(
        self,
        message: str,
        module_id: str,
        slide_id: str,
        cache_key: str,
        preload: bool,
    ) -> None:
        if (
            module_id == self._preview_module_id
            and slide_id == self._preview_slide_id
            and cache_key == self._preview_cache_key
        ):
            self.slide_preview.set_error(message)
            self._log(f"当前页预览失败：{message}")

    def _on_preview_finished(self) -> None:
        worker = self._preview_worker
        was_preload = bool(worker and worker.preload)
        if worker:
            worker.deleteLater()
        self._preview_worker = None
        if self._preview_pending:
            self._preview_pending = False
            self._preview_timer.start(0)
            return
        if not was_preload and self._preload_after_current:
            module_id, slide_id = self._preload_after_current
            self._preload_after_current = None
            self._queue_adjacent_previews(module_id, slide_id)
        self._start_next_preload()

    def _queue_adjacent_previews(self, module_id: str, slide_id: str) -> None:
        if not self.manifest:
            return
        contexts = rebuild_structure_context(self.project, self.manifest)
        current_index = next(
            (
                index
                for index, context in enumerate(contexts)
                if context.module_id == module_id and context.slide_id == slide_id
            ),
            -1,
        )
        if current_index < 0:
            return
        neighbor_indexes = [current_index + 1, current_index - 1]
        queued_keys = {item[2] for item in self._preload_queue}
        for index in neighbor_indexes:
            if not 0 <= index < len(contexts):
                continue
            context = contexts[index]
            try:
                cache_key = preview_fingerprint(
                    self.project,
                    context.module_id,
                    context.slide_id,
                )
            except Exception:
                continue
            if cache_key in self._preview_cache or cache_key in queued_keys:
                continue
            self._preload_queue.append(
                (context.module_id, context.slide_id, cache_key)
            )
            queued_keys.add(cache_key)

    def _start_next_preload(self) -> None:
        if (
            self._preview_worker
            or not self._office_preview_session
            or not self._office_preview_enabled
        ):
            return
        while self._preload_queue:
            module_id, slide_id, cache_key = self._preload_queue.pop(0)
            if cache_key in self._preview_cache:
                continue
            worker = PreviewWorker(
                project=self.project,
                module_id=module_id,
                slide_id=slide_id,
                cache_key=cache_key,
                office_session=self._office_preview_session,
                preload=True,
                parent=self,
            )
            worker.succeeded.connect(self._on_preview_succeeded)
            worker.failed.connect(self._on_preview_failed)
            worker.finished.connect(self._on_preview_finished)
            self._preview_worker = worker
            worker.start()
            return

    def _refresh_structure(self) -> None:
        self.structure_list.clear()
        position = 1
        for module in self.project.modules:
            if not module.enabled or not module.slides:
                continue
            page_count = len(module.slides)
            detail = "、".join(
                (slide.title or f"第{index}页")
                for index, slide in enumerate(module.slides, start=1)
            )
            if len(detail) > 46:
                detail = detail[:43] + "…"
            item = QListWidgetItem(
                f"{position:02d}   {module.name}\n       {page_count} 页 · {detail}"
            )
            item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
            self.structure_list.addItem(item)
            position += 1
        if position == 1:
            self.structure_list.addItem("尚未启用任何模块")

    # ----- sources ---------------------------------------------------------

    def import_excel_for_mapping(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择用于映射的 Excel",
            str(PROJECT_ROOT),
            "Excel 工作簿 (*.xlsx *.xlsm)",
        )
        if path:
            self._load_excel_path(path)

    def _load_excel_path(self, path: str, *, log: bool = True) -> None:
        try:
            preview = load_excel_preview(path)
            self.excel_preview = preview
            self.project.excel_path = str(preview.path)
            self.excel_path_edit.setText(str(preview.path))
            self.excel_sheet_combo.blockSignals(True)
            self.excel_sheet_combo.clear()
            self.excel_sheet_combo.addItems([sheet.name for sheet in preview.sheets])
            self.excel_sheet_combo.blockSignals(False)
            if preview.sheets:
                self.excel_sheet_combo.setCurrentIndex(0)
                self._show_excel_sheet(preview.sheets[0].name)
            existing = next(
                (record for record in self.project.sources if record.path == str(preview.path)),
                None,
            )
            if existing:
                existing.kind = "Excel"
            else:
                self.project.sources.append(SourceRecord(path=str(preview.path), kind="Excel"))
            self._refresh_sources()
            self._refresh_excel_mapping_table()
            self._refresh_summary()
            if log:
                self._log(f"Excel 工作簿已加载：{preview.path}（{len(preview.sheets)} 个工作表）")
        except Exception as exc:
            self._error("Excel 加载失败", str(exc))

    def _restore_excel_state(self) -> None:
        path = Path(self.project.excel_path) if self.project.excel_path else None
        if path and path.is_file():
            self._load_excel_path(str(path), log=False)
            return
        self.excel_preview = None
        self.excel_path_edit.clear()
        self.excel_sheet_combo.clear()
        self.excel_preview_table.clear()
        self.excel_preview_table.setRowCount(0)
        self.excel_preview_table.setColumnCount(0)
        self._refresh_excel_mapping_table()

    def _show_excel_sheet(self, name: str) -> None:
        if not self.excel_preview or not name:
            return
        try:
            sheet = self.excel_preview.sheet(name)
        except ValueError:
            return
        self.excel_preview_table.setUpdatesEnabled(False)
        try:
            self.excel_preview_table.clear()
            self.excel_preview_table.setRowCount(sheet.row_count)
            self.excel_preview_table.setColumnCount(sheet.column_count)
            self.excel_preview_table.setHorizontalHeaderLabels(
                [get_column_letter(index) for index in range(1, sheet.column_count + 1)]
            )
            self.excel_preview_table.setVerticalHeaderLabels(
                [str(index) for index in range(1, sheet.row_count + 1)]
            )
            for row_index, row in enumerate(sheet.values):
                for column_index in range(sheet.column_count):
                    value = row[column_index] if column_index < len(row) else ""
                    item = QTableWidgetItem(value)
                    item.setToolTip(value)
                    self.excel_preview_table.setItem(row_index, column_index, item)
            self.excel_preview_table.resizeColumnsToContents()
            for column in range(sheet.column_count):
                self.excel_preview_table.setColumnWidth(
                    column, min(220, max(72, self.excel_preview_table.columnWidth(column)))
                )
        finally:
            self.excel_preview_table.setUpdatesEnabled(True)
        self.excel_header_spin.setRange(1, max(1, sheet.row_count))
        self.excel_header_spin.setValue(detect_header_row(sheet))
        self._highlight_excel_header()
        self._refresh_excel_mapping_table()
        if sheet.truncated:
            self._log(
                f"工作表 {name} 较大，界面只预览前 {sheet.row_count} 行、{sheet.column_count} 列。"
            )

    def _highlight_excel_header(self, *_args) -> None:
        if not hasattr(self, "excel_preview_table"):
            return
        header_row = self.excel_header_spin.value() - 1
        for row in range(self.excel_preview_table.rowCount()):
            background = QColor("#e8f1ff") if row == header_row else QColor("#ffffff")
            for column in range(self.excel_preview_table.columnCount()):
                item = self.excel_preview_table.item(row, column)
                if item:
                    item.setBackground(background)

    def detect_excel_header(self) -> None:
        sheet = self._current_excel_sheet()
        if not sheet:
            self._error("无法识别表头", "请先选择 Excel 和工作表。")
            return
        row = detect_header_row(sheet)
        self.excel_header_spin.setValue(row)
        self._log(f"表头建议：{sheet.name} 第 {row} 行；可手工修改。")

    def _current_excel_sheet(self):
        if not self.excel_preview:
            return None
        name = self.excel_sheet_combo.currentText()
        try:
            return self.excel_preview.sheet(name)
        except ValueError:
            return None

    def _selected_excel_range(self) -> str | None:
        ranges = self.excel_preview_table.selectedRanges()
        if len(ranges) != 1:
            return None
        selected = ranges[0]
        return selection_range(
            selected.topRow() + 1,
            selected.leftColumn() + 1,
            selected.bottomRow() + 1,
            selected.rightColumn() + 1,
        )

    def _update_excel_selection_label(self) -> None:
        source_range = self._selected_excel_range()
        if not source_range:
            self.excel_selection_label.setText("请选择一个连续的单元格或矩形区域。")
            return
        rows, columns = range_shape(source_range)
        self.excel_selection_label.setText(
            f"当前选区：{self.excel_sheet_combo.currentText()}!{source_range}    "
            f"{rows} 行 × {columns} 列"
        )

    def _mapping_target_slots(self) -> list[dict[str, Any]]:
        if not self.manifest:
            return []
        return [slot for slot in self.manifest.slots if slot["kind"] in {"text", "table"}]

    def add_excel_mapping(self) -> None:
        sheet = self._current_excel_sheet()
        source_range = self._selected_excel_range()
        if not sheet or not source_range:
            self._error("无法添加映射", "请先在 Excel 预览中选择一个连续区域。")
            return
        selected_shape = range_shape(source_range)
        choices: list[tuple[str, dict[str, Any]]] = []
        for slot in self._mapping_target_slots():
            if slot["kind"] == "table":
                expected = (slot.get("rows"), slot.get("columns"))
                if selected_shape != expected:
                    continue
                detail = f"表格 {expected[0]}×{expected[1]}"
            else:
                detail = "文字"
            choices.append(
                (f"{self._slot_label(slot)} · 第{slot['slide']}页 · {detail}", slot)
            )
        if not choices:
            self._error(
                "没有兼容目标",
                f"当前选区为 {selected_shape[0]}×{selected_shape[1]}。"
                "文字字段可接收任意范围；表格字段必须与模板尺寸一致。",
            )
            return
        display_items = [item[0] for item in choices]
        selected_label, accepted = QInputDialog.getItem(
            self, "选择 PPT 目标", "将 Excel 选区写入：", display_items, 0, False
        )
        if not accepted:
            return
        slot = dict(choices[display_items.index(selected_label)][1])
        existing = [
            rule for rule in self.project.excel_mappings if rule.target_slot == slot["key"]
        ]
        if existing:
            answer = QMessageBox.question(
                self,
                "替换已有映射",
                f"{self._slot_label(slot)} 已有映射。是否替换为 {sheet.name}!{source_range}？",
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
            self.project.excel_mappings = [
                rule for rule in self.project.excel_mappings if rule.target_slot != slot["key"]
            ]
        self.project.excel_mappings.append(
            ExcelMappingRule(
                sheet=sheet.name,
                source_range=source_range,
                target_slot=slot["key"],
                mode="table" if slot["kind"] == "table" else "text",
            )
        )
        self._refresh_excel_mapping_table()
        self._log(f"已添加映射：{sheet.name}!{source_range} → {self._slot_label(slot)}")

    def suggest_excel_mappings(self) -> None:
        sheet = self._current_excel_sheet()
        if not sheet:
            self._error("无法自动建议", "请先选择 Excel 和工作表。")
            return
        targets = [
            MappingTarget(slot["key"], self._slot_label(slot), slot["kind"])
            for slot in self._mapping_target_slots()
        ]
        existing_targets = {rule.target_slot for rule in self.project.excel_mappings}
        suggestions = [
            rule
            for rule in suggest_text_mappings(sheet, targets)
            if rule.target_slot not in existing_targets
        ]
        for rule in suggestions:
            rule.enabled = False
        self.project.excel_mappings.extend(suggestions)
        self._refresh_excel_mapping_table()
        if suggestions:
            self._log(f"发现 {len(suggestions)} 条字段建议；默认未启用，请勾选确认。")
            QMessageBox.information(
                self,
                "字段建议",
                f"找到 {len(suggestions)} 条候选映射。为避免误填，建议默认未启用，请逐条确认。",
            )
        else:
            QMessageBox.information(self, "字段建议", "当前工作表没有找到可信度足够的字段建议。")

    def _refresh_excel_mapping_table(self) -> None:
        if not hasattr(self, "excel_mapping_table"):
            return
        self._excel_mapping_guard = True
        try:
            rules = self.project.excel_mappings
            self.excel_mapping_table.setRowCount(len(rules))
            for row, rule in enumerate(rules):
                enabled = QTableWidgetItem("")
                enabled.setFlags(
                    (enabled.flags() | Qt.ItemFlag.ItemIsUserCheckable)
                    & ~Qt.ItemFlag.ItemIsEditable
                )
                enabled.setCheckState(
                    Qt.CheckState.Checked if rule.enabled else Qt.CheckState.Unchecked
                )
                sheet_item = QTableWidgetItem(rule.sheet)
                range_item = QTableWidgetItem(rule.source_range)
                slot = self._slot_by_key(rule.target_slot)
                target_item = QTableWidgetItem(
                    self._slot_label(slot) if slot else rule.target_slot
                )
                mode_item = QTableWidgetItem("表格" if rule.mode == "table" else "文字")
                preview_text = ""
                if self.excel_preview:
                    try:
                        preview_value = preview_rule_value(
                            self.excel_preview.sheet(rule.sheet), rule
                        )
                        if isinstance(preview_value, list):
                            rows = len(preview_value)
                            columns = len(preview_value[0]) if preview_value else 0
                            first = " | ".join(preview_value[0]) if preview_value else ""
                            preview_text = f"{rows}×{columns} · {first}"
                        else:
                            preview_text = str(preview_value).replace("\n", " | ")
                    except Exception as exc:
                        preview_text = f"无效：{exc}"
                preview_item = QTableWidgetItem(preview_text)
                for item in (sheet_item, range_item, target_item, mode_item, preview_item):
                    item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.excel_mapping_table.setItem(row, 0, enabled)
                self.excel_mapping_table.setItem(row, 1, sheet_item)
                self.excel_mapping_table.setItem(row, 2, range_item)
                self.excel_mapping_table.setItem(row, 3, target_item)
                self.excel_mapping_table.setItem(row, 4, mode_item)
                self.excel_mapping_table.setItem(row, 5, preview_item)
                self.excel_mapping_table.setRowHeight(row, 34)
        finally:
            self._excel_mapping_guard = False

    def _on_excel_mapping_item_changed(self, item: QTableWidgetItem) -> None:
        if self._excel_mapping_guard or item.column() != 0:
            return
        if 0 <= item.row() < len(self.project.excel_mappings):
            self.project.excel_mappings[item.row()].enabled = (
                item.checkState() == Qt.CheckState.Checked
            )

    def remove_excel_mapping(self) -> None:
        row = self.excel_mapping_table.currentRow()
        if 0 <= row < len(self.project.excel_mappings):
            del self.project.excel_mappings[row]
            self._refresh_excel_mapping_table()

    def apply_excel_mapping_rules(self) -> None:
        if not self.project.excel_path:
            self._error("无法应用映射", "请先选择 Excel 文件。")
            return
        if not any(rule.enabled for rule in self.project.excel_mappings):
            self._error("无法应用映射", "至少需要启用一条映射规则。")
            return
        slot_specs = {slot["key"]: slot for slot in self.manifest.slots} if self.manifest else {}
        try:
            result = apply_excel_mappings(
                self.project.excel_path,
                self.project.excel_mappings,
                slot_specs=slot_specs,
            )
            self.project.values.update(result.values)
            self._refresh_field_table()
            self._refresh_summary()
            self._log(f"Excel 映射已应用：更新 {len(result.values)} 个 PPT 字段。")
            if result.warnings:
                QMessageBox.warning(self, "映射警告", "\n".join(result.warnings))
            else:
                QMessageBox.information(
                    self, "映射完成", f"已更新 {len(result.values)} 个 PPT 字段。"
                )
        except Exception as exc:
            self._error("Excel 映射失败", str(exc))

    def generate_inspection_from_far(self) -> None:
        if not self.manifest:
            self._error("无法生成检测效果", "请先加载 PPT 模板和模板配置。")
            return
        path = Path(self.project.excel_path) if self.project.excel_path else None
        if not path or not path.is_file():
            selected, _ = QFileDialog.getOpenFileName(
                self,
                "选择光学 FAR",
                str(PROJECT_ROOT / "templates" / "光学资料"),
                "光学 FAR (*.xlsx)",
            )
            if not selected:
                return
            path = Path(selected)
            self._load_excel_path(str(path))

        self.far_generate_button.setEnabled(False)
        self.far_generate_button.setText("正在解析 FAR…")
        QApplication.processEvents()
        try:
            data = parse_optical_far(path)
            result = apply_optical_far(
                self.project,
                self.manifest,
                data,
                FAR_ASSET_ROOT,
            )
            self._preview_timer.stop()
            self._preview_cache.clear()
            self._preload_queue.clear()
            self._preload_after_current = None
            self.module_editor.set_state(self.project, self.manifest)
            self._refresh_structure()
            self._refresh_field_table()
            self._refresh_source_targets()
            self._refresh_sources()
            self._refresh_asset_table()
            self._refresh_summary()
            self.tabs.setCurrentIndex(1)
            message = (
                f"已生成 {result.effect_pages} 页检测效果，"
                f"写入 {result.requirements} 条检测标准和 {result.stations} 个工位配置。\n\n"
                f"相机：{result.camera_summary or '未读取'}\n"
                f"镜头：{result.lens_summary or '未读取'}\n"
                f"光源：{result.light_summary or '未读取'}\n\n"
                "请在“方案模块”审核页面，确认后点击“生成 PPT”。"
            )
            self._log(
                f"光学 FAR 已应用：{path.name} · "
                f"{result.stations}工位 · {result.requirements}检测项 · "
                f"{result.effect_pages}页检测效果"
            )
            QMessageBox.information(self, "检测效果生成完成", message)
        except Exception as exc:
            self._error("光学 FAR 生成失败", str(exc))
        finally:
            self.far_generate_button.setEnabled(True)
            self.far_generate_button.setText("一键生成检测效果")

    def save_excel_mappings(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self,
            "保存 Excel 映射规则",
            str(PROJECT_ROOT / "examples" / "excel_mapping.json"),
            "Excel 映射规则 (*.json)",
        )
        if not path:
            return
        try:
            payload = {
                "schema_version": 1,
                "excel_filename": Path(self.project.excel_path).name if self.project.excel_path else "",
                "mappings": [rule.to_dict() for rule in self.project.excel_mappings],
            }
            Path(path).write_text(
                json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            self._log(f"Excel 映射规则已保存：{path}")
        except Exception as exc:
            self._error("保存映射失败", str(exc))

    def load_excel_mappings(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "加载 Excel 映射规则",
            str(PROJECT_ROOT / "examples"),
            "Excel 映射规则 (*.json)",
        )
        if not path:
            return
        try:
            raw = json.loads(Path(path).read_text(encoding="utf-8"))
            if raw.get("schema_version") != 1 or not isinstance(raw.get("mappings"), list):
                raise ValueError("映射规则格式无效")
            rules = [ExcelMappingRule.from_dict(item) for item in raw["mappings"]]
            known = {slot["key"] for slot in self.manifest.slots} if self.manifest else set()
            unknown = sorted({rule.target_slot for rule in rules} - known)
            if unknown:
                raise ValueError(f"当前 PPT 模板不存在目标：{', '.join(unknown)}")
            self.project.excel_mappings = rules
            self._refresh_excel_mapping_table()
            self._log(f"Excel 映射规则已加载：{path}")
        except Exception as exc:
            self._error("加载映射失败", str(exc))

    def import_sources(self) -> None:
        paths, _ = QFileDialog.getOpenFileNames(
            self,
            "导入 Excel / Word 资料",
            str(PROJECT_ROOT),
            "资料文件 (*.xlsx *.xlsm *.docx)",
        )
        existing = {record.path for record in self.project.sources}
        for path in paths:
            if path not in existing:
                self.project.sources.append(SourceRecord(path=path))
        if paths:
            self._refresh_sources()
            self.source_list.setCurrentRow(len(self.project.sources) - 1)
            self._refresh_summary()

    def _refresh_sources(self) -> None:
        current = self.source_list.currentRow()
        self.source_list.clear()
        for record in self.project.sources:
            status = record.kind or "待解析"
            self.source_list.addItem(f"{Path(record.path).name}    ·    {status}")
        if self.project.sources:
            self.source_list.setCurrentRow(min(max(current, 0), len(self.project.sources) - 1))

    def _show_source_preview(self, row: int) -> None:
        if 0 <= row < len(self.project.sources):
            self.source_preview.setPlainText(self.project.sources[row].content)
        else:
            self.source_preview.clear()

    def parse_selected_source(self) -> None:
        row = self.source_list.currentRow()
        if row < 0:
            QMessageBox.information(self, "解析资料", "请先选择一个资料文件。")
            return
        record = self.project.sources[row]
        try:
            parsed = parse_source(record.path)
            record.kind = parsed.kind
            record.content = parsed.content
            self._refresh_sources()
            self.source_list.setCurrentRow(row)
            suffix = "，预览已截断" if parsed.truncated else ""
            self._log(f"已解析 {parsed.kind}：{record.path}（{parsed.section_count} 个内容区{suffix}）")
        except Exception as exc:
            self._error("资料解析失败", f"{record.path}\n\n{exc}")

    def remove_selected_source(self) -> None:
        row = self.source_list.currentRow()
        if 0 <= row < len(self.project.sources):
            del self.project.sources[row]
            self._refresh_sources()
            self._refresh_summary()

    def _refresh_source_targets(self) -> None:
        self.source_target_combo.clear()
        if not self.manifest:
            return
        for slot in self.manifest.slots:
            if slot["kind"] == "text":
                self.source_target_combo.addItem(self._slot_label(slot), slot["key"])

    def apply_source_to_slot(self) -> None:
        key = self.source_target_combo.currentData()
        if not key:
            return
        cursor = self.source_preview.textCursor()
        content = cursor.selectedText().replace("\u2029", "\n") if cursor.hasSelection() else self.source_preview.toPlainText()
        if not content.strip():
            QMessageBox.information(self, "应用资料", "解析预览中没有可应用的内容。")
            return
        slot = self._slot_by_key(str(key))
        max_chars = slot.get("max_chars") if slot else None
        if max_chars and len(content) > max_chars:
            answer = QMessageBox.question(
                self,
                "内容可能过长",
                f"选中内容有 {len(content)} 字，模板建议不超过 {max_chars} 字。仍然填入后再编辑吗？",
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
        self.project.values[str(key)] = content
        self._refresh_field_table()
        self._refresh_summary()
        self.tabs.setCurrentIndex(0)
        self._log(f"已将资料内容填入：{self._slot_label(slot or {'key': key})}")

    # ----- assets ----------------------------------------------------------

    def import_images(self) -> None:
        paths, _ = QFileDialog.getOpenFileNames(
            self,
            "导入图片素材",
            str(PROJECT_ROOT / "assets"),
            "图片 (*.png *.jpg *.jpeg *.bmp *.webp)",
        )
        existing = {asset.path for asset in self.project.assets}
        for path in paths:
            if path not in existing:
                self.project.assets.append(AssetRecord(path=path))
        if paths:
            self._refresh_asset_table()
            self._refresh_summary()

    def _image_slots(self) -> list[dict[str, Any]]:
        if not self.manifest:
            return []
        return [slot for slot in self.manifest.slots if slot["kind"] == "image"]

    def _refresh_asset_table(self) -> None:
        self._asset_guard = True
        try:
            self.asset_table.setRowCount(len(self.project.assets))
            image_slots = self._image_slots()
            for row, asset in enumerate(self.project.assets):
                preview_item = QTableWidgetItem()
                pixmap = QPixmap(asset.path)
                if not pixmap.isNull():
                    preview_item.setIcon(QIcon(pixmap.scaled(64, 48, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)))
                preview_item.setFlags(preview_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                file_item = QTableWidgetItem(Path(asset.path).name)
                file_item.setToolTip(asset.path)
                file_item.setFlags(file_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                category_combo = QComboBox()
                category_combo.addItems(ASSET_CATEGORIES)
                category_combo.setCurrentText(asset.category if asset.category in ASSET_CATEGORIES else "未分类")
                category_combo.currentTextChanged.connect(
                    lambda value, record=asset: self._set_asset_category(record, value)
                )
                slot_combo = QComboBox()
                slot_combo.addItem("仅作为素材", "")
                for slot in image_slots:
                    slot_combo.addItem(self._slot_label(slot), slot["key"])
                slot_index = slot_combo.findData(asset.slot_key)
                slot_combo.setCurrentIndex(max(0, slot_index))
                slot_combo.currentIndexChanged.connect(
                    lambda _index, combo=slot_combo, record=asset: self._set_asset_slot(
                        record, str(combo.currentData() or "")
                    )
                )
                self.asset_table.setItem(row, 0, preview_item)
                self.asset_table.setItem(row, 1, file_item)
                self.asset_table.setCellWidget(row, 2, category_combo)
                self.asset_table.setCellWidget(row, 3, slot_combo)
                self.asset_table.setRowHeight(row, 58)
        finally:
            self._asset_guard = False

    def _set_asset_category(self, asset: AssetRecord, category: str) -> None:
        if not self._asset_guard:
            asset.category = category

    def _set_asset_slot(self, asset: AssetRecord, slot_key: str) -> None:
        if self._asset_guard:
            return
        old_slot = asset.slot_key
        if old_slot and self.project.values.get(old_slot) == asset.path:
            self.project.values.pop(old_slot, None)
        if slot_key:
            for other in self.project.assets:
                if other is not asset and other.slot_key == slot_key:
                    other.slot_key = ""
            self.project.values[slot_key] = asset.path
        asset.slot_key = slot_key
        self._refresh_asset_table()
        self._refresh_field_table()
        self._refresh_summary()

    def remove_selected_asset(self) -> None:
        row = self.asset_table.currentRow()
        if not 0 <= row < len(self.project.assets):
            return
        asset = self.project.assets.pop(row)
        if asset.slot_key and self.project.values.get(asset.slot_key) == asset.path:
            self.project.values.pop(asset.slot_key, None)
        self._refresh_asset_table()
        self._refresh_field_table()
        self._refresh_summary()

    # ----- generation ------------------------------------------------------

    def _refresh_summary(self, *_args) -> None:
        if not hasattr(self, "metric_slides"):
            return
        pages = sum(
            len(module.slides)
            for module in self.project.modules
            if module.enabled
        )
        enabled_modules = sum(module.enabled for module in self.project.modules)
        filled = sum(value not in (None, "", []) for value in self.project.values.values())
        filled += sum(
            value not in (None, "", [])
            for module in self.project.modules
            for value in module.module_values.values()
        )
        filled += sum(
            value not in (None, "", [])
            for module in self.project.modules
            for slide in module.slides
            for value in slide.overrides.values()
        )
        self.metric_slides.setText(str(pages))
        self.metric_modules.setText(str(enabled_modules))
        self.metric_fields.setText(str(filled))
        self.metric_sources.setText(str(len(self.project.sources)))
        self.metric_assets.setText(str(len(self.project.assets)))

    def _required_missing(self) -> list[str]:
        if not self.manifest:
            return ["模板配置"]
        missing: list[str] = []
        for context in rebuild_structure_context(self.project, self.manifest):
            for slot in slot_specs_for_source_slide(
                self.manifest, context.source_slide
            ):
                value = context.values.get(slot["key"])
                if not slot.get("required"):
                    continue
                if value is None or (
                    slot["kind"] == "text" and not str(value).strip()
                ):
                    missing.append(
                        f"{context.module_title} / {context.slide_title} / {self._slot_label(slot)}"
                    )
        return missing

    def generate_presentation(self) -> None:
        if self._worker and self._worker.isRunning():
            return
        self._collect_ui_state()
        if not any(module.enabled and module.slides for module in self.project.modules):
            self._error("无法生成", "至少需要启用一个方案模块。")
            return
        missing = self._required_missing()
        if missing:
            self._error("缺少必填内容", "请先填写：" + "、".join(missing))
            return
        if not self.project.output_path:
            self.choose_output()
        if not self.project.output_path:
            return
        output = Path(self.project.output_path)
        overwrite = False
        if output.exists():
            answer = QMessageBox.question(
                self,
                "确认覆盖",
                f"输出文件已经存在：\n{output}\n\n是否覆盖？",
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
            overwrite = True

        self._worker = RenderWorker(
            project=self.project,
            overwrite=overwrite,
            parent=self,
        )
        self._worker.succeeded.connect(self._on_render_succeeded)
        self._worker.failed.connect(self._on_render_failed)
        self._worker.finished.connect(self._on_render_finished)
        self.generate_button.setEnabled(False)
        self.generate_button.setText("正在生成…")
        self.progress.show()
        self.statusBar().showMessage("正在生成 PPT，请稍候…")
        self._log(f"开始生成：{self.project.output_path}")
        self._worker.start()

    def _on_render_succeeded(self, output: str) -> None:
        self.project.output_path = output
        self.output_path_edit.setText(output)
        self.open_result_button.setEnabled(True)
        self._log(f"生成完成：{output}")
        self.statusBar().showMessage("PPT 生成完成", 6000)
        QMessageBox.information(self, "生成完成", f"PPT 已生成：\n{output}")

    def _on_render_failed(self, message: str) -> None:
        self._log(f"生成失败：{message}")
        self._error("PPT 生成失败", message)

    def _on_render_finished(self) -> None:
        self.generate_button.setEnabled(True)
        self.generate_button.setText("生成 PPT")
        self.progress.hide()
        if self._worker:
            self._worker.deleteLater()
        self._worker = None

    def open_output(self) -> None:
        path = Path(self.project.output_path)
        if not path.is_file():
            self._error("无法打开", "尚未生成输出文件。")
            return
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(path)))

    # ----- helpers ---------------------------------------------------------

    def _log(self, message: str) -> None:
        self.log_view.appendPlainText(message)

    def _error(self, title: str, message: str) -> None:
        self.statusBar().showMessage(title, 5000)
        QMessageBox.critical(self, title, message)

    def closeEvent(self, event) -> None:
        if self._worker and self._worker.isRunning():
            QMessageBox.warning(self, "正在生成", "请等待 PPT 生成完成后再关闭软件。")
            event.ignore()
            return
        self._preview_timer.stop()
        if self._thumbnail_worker and self._thumbnail_worker.isRunning():
            self._thumbnail_worker.cancel()
            if not self._thumbnail_worker.wait(3000):
                QMessageBox.warning(
                    self,
                    "正在停止模板预览",
                    "模板缩略图任务仍在结束，请稍后再关闭软件。",
                )
                event.ignore()
                return
        if self._preview_worker and self._preview_worker.isRunning():
            self._preview_worker.cancel()
            if not self._preview_worker.wait(3000):
                QMessageBox.warning(
                    self,
                    "正在停止预览",
                    "当前页预览仍在结束，请稍后再关闭软件。",
                )
                event.ignore()
                return
        if self._office_preview_session:
            self._office_preview_session.close()
        super().closeEvent(event)
