"""Module tree and structured page editor for schema-v2 PPT projects."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl.utils.cell import get_column_letter
from PySide6.QtCore import Qt, Signal
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QInputDialog,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
)

from ..excel_mapper import ExcelWorkbookPreview, load_excel_preview
from ..module_service import (
    add_module,
    add_page_template,
    add_slide,
    duplicate_module,
    duplicate_slide,
    materialize_excel_modules,
    module_by_id,
    move_module,
    move_slide,
    page_template_by_key,
    read_excel_records,
    remove_module,
    remove_slide,
    set_default_page_template,
    slide_by_id,
    slot_specs_for_source_slide,
    sync_legacy_module_state,
)
from ..project import ExcelModuleBinding, PptProject, ProjectModule, ProjectSlide
from ..template_renderer import TemplateManifest
from .dialogs import TableValueDialog, TextValueDialog


MODULE_TYPE_NAMES = {
    "fixed": "固定模块",
    "repeat": "可重复模块",
    "dataDriven": "数据驱动模块",
}
SLOT_KIND_NAMES = {"text": "文字", "image": "图片", "table": "表格"}


def _normalized(value: str) -> str:
    return "".join(character.lower() for character in value if character.isalnum())


class ExcelModuleBindingDialog(QDialog):
    """Small column-to-module binding dialog for repeated module generation."""

    def __init__(
        self,
        project: PptProject,
        module: ProjectModule,
        manifest: TemplateManifest,
        slot_labels: dict[str, str] | None = None,
        parent=None,
    ):
        super().__init__(parent)
        self.project = project
        self.module = module
        self.manifest = manifest
        self.slot_labels = slot_labels or {}
        self.preview: ExcelWorkbookPreview | None = None
        self.existing = next(
            (
                item
                for item in project.module_bindings
                if item.source_module_id == module.id
            ),
            None,
        )
        self.setWindowTitle(f"Excel 动态生成模块 · {module.name}")
        self.resize(760, 620)
        layout = QVBoxLayout(self)

        source_group = QGroupBox("Excel 数据源")
        source_form = QFormLayout(source_group)
        path_row = QHBoxLayout()
        self.path_edit = QLineEdit()
        self.path_edit.setReadOnly(True)
        path_row.addWidget(self.path_edit, 1)
        choose_button = QPushButton("选择 Excel")
        choose_button.clicked.connect(self._choose_excel)
        path_row.addWidget(choose_button)
        source_form.addRow("文件", path_row)
        self.sheet_combo = QComboBox()
        self.sheet_combo.currentTextChanged.connect(self._refresh_headers)
        source_form.addRow("工作表", self.sheet_combo)
        self.header_spin = QSpinBox()
        self.header_spin.setRange(1, 1)
        self.header_spin.valueChanged.connect(self._refresh_headers)
        source_form.addRow("表头行", self.header_spin)
        self.range_edit = QLineEdit()
        self.range_edit.setPlaceholderText("可留空，表示表头后所有数据；例如 A2:I20")
        source_form.addRow("数据范围", self.range_edit)
        layout.addWidget(source_group)

        layout.addWidget(QLabel("把 Excel 列映射到该模块页面中的结构化字段："))
        self.mapping_table = QTableWidget(0, 2)
        self.mapping_table.setHorizontalHeaderLabels(["Excel 列", "模块字段"])
        self.mapping_table.verticalHeader().setVisible(False)
        self.mapping_table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.ResizeMode.Stretch
        )
        self.mapping_table.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.ResizeMode.Stretch
        )
        layout.addWidget(self.mapping_table, 1)

        footer_form = QFormLayout()
        self.name_field_combo = QComboBox()
        self.name_field_combo.addItem("按序号命名", "")
        footer_form.addRow("模块名称来源", self.name_field_combo)
        self.summary_label = QLabel("选择 Excel 后显示可生成记录数")
        self.summary_label.setObjectName("mutedLabel")
        footer_form.addRow("预计结果", self.summary_label)
        layout.addLayout(footer_form)

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(self._validate_and_accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

        initial_path = (
            self.existing.source_path
            if self.existing
            else project.excel_path
        )
        if initial_path and Path(initial_path).is_file():
            self._load_excel(initial_path)

    def _targets(self) -> list[tuple[str, str]]:
        source_slides = {
            template.source_slide for template in self.module.page_templates
        }
        targets: list[tuple[str, str]] = []
        seen: set[str] = set()
        for slot in self.manifest.slots:
            if slot["slide"] not in source_slides or slot["key"] in seen:
                continue
            seen.add(slot["key"])
            label = str(
                slot.get("label")
                or self.slot_labels.get(slot["key"])
                or slot["key"]
            )
            targets.append((f"{label} · {{{{{slot['key']}}}}}", slot["key"]))
        targets.extend(
            [
                ("页面标题 · {{slide_title}}", "slide_title"),
                ("页面小标题 · {{slide_subtitle}}", "slide_subtitle"),
            ]
        )
        return targets

    def _choose_excel(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择模块数据 Excel",
            str(Path(self.project.excel_path).parent if self.project.excel_path else Path.cwd()),
            "Excel 工作簿 (*.xlsx *.xlsm)",
        )
        if path:
            self._load_excel(path)

    def _load_excel(self, path: str) -> None:
        try:
            self.preview = load_excel_preview(path)
        except Exception as exc:
            QMessageBox.critical(self, "Excel 加载失败", str(exc))
            return
        self.path_edit.setText(str(self.preview.path))
        self.sheet_combo.blockSignals(True)
        self.sheet_combo.clear()
        self.sheet_combo.addItems([item.name for item in self.preview.sheets])
        if self.existing and self.existing.sheet:
            self.sheet_combo.setCurrentText(self.existing.sheet)
        self.sheet_combo.blockSignals(False)
        maximum = max((item.row_count for item in self.preview.sheets), default=1)
        self.header_spin.setRange(1, max(1, maximum))
        self.header_spin.setValue(self.existing.header_row if self.existing else 1)
        self.range_edit.setText(self.existing.data_range if self.existing else "")
        self._refresh_headers()

    def _headers(self) -> list[str]:
        if not self.preview or not self.sheet_combo.currentText():
            return []
        sheet = self.preview.sheet(self.sheet_combo.currentText())
        index = self.header_spin.value() - 1
        if index < 0 or index >= sheet.row_count:
            return []
        return [value.strip() for value in sheet.values[index] if value.strip()]

    def _refresh_headers(self, *_args) -> None:
        headers = self._headers()
        targets = self._targets()
        existing_map = self.existing.field_map if self.existing else {}
        self.mapping_table.setRowCount(len(headers))
        target_norms = {
            key: {_normalized(key), _normalized(label.split("·", 1)[0])}
            for label, key in targets
        }
        for row, header in enumerate(headers):
            source_item = QTableWidgetItem(header)
            source_item.setFlags(source_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
            self.mapping_table.setItem(row, 0, source_item)
            combo = QComboBox()
            combo.addItem("不映射", "")
            for label, key in targets:
                combo.addItem(label, key)
            requested = existing_map.get(header, "")
            if not requested:
                header_norm = _normalized(header)
                requested = next(
                    (
                        key
                        for key, candidates in target_norms.items()
                        if header_norm and header_norm in candidates
                    ),
                    "",
                )
            if requested:
                index = combo.findData(requested)
                if index >= 0:
                    combo.setCurrentIndex(index)
            self.mapping_table.setCellWidget(row, 1, combo)

        self.name_field_combo.clear()
        self.name_field_combo.addItem("按序号命名", "")
        for header in headers:
            self.name_field_combo.addItem(header, header)
        name_field = self.existing.module_name_field if self.existing else ""
        if name_field:
            index = self.name_field_combo.findData(name_field)
            if index >= 0:
                self.name_field_combo.setCurrentIndex(index)
        self._refresh_summary()

    def _draft_binding(self) -> ExcelModuleBinding:
        field_map: dict[str, str] = {}
        for row in range(self.mapping_table.rowCount()):
            source_item = self.mapping_table.item(row, 0)
            combo = self.mapping_table.cellWidget(row, 1)
            if source_item and isinstance(combo, QComboBox) and combo.currentData():
                field_map[source_item.text()] = str(combo.currentData())
        arguments: dict[str, Any] = {
            "source_module_id": self.module.id,
            "source_path": self.path_edit.text(),
            "sheet": self.sheet_combo.currentText(),
            "header_row": self.header_spin.value(),
            "data_range": self.range_edit.text().strip().upper(),
            "field_map": field_map,
            "module_name_field": str(self.name_field_combo.currentData() or ""),
        }
        if self.existing:
            arguments["id"] = self.existing.id
        return ExcelModuleBinding(**arguments)

    def _refresh_summary(self) -> None:
        if not self.preview or not self.sheet_combo.currentText():
            self.summary_label.setText("选择 Excel 后显示可生成记录数")
            return
        try:
            count = len(read_excel_records(self._draft_binding()))
            pages = count * len(self.module.slides)
            self.summary_label.setText(
                f"{count} 条记录 × 当前模块 {len(self.module.slides)} 页 = {pages} 页"
            )
        except Exception as exc:
            self.summary_label.setText(str(exc))

    def _validate_and_accept(self) -> None:
        try:
            binding = self._draft_binding()
            if not binding.field_map:
                raise ValueError("至少映射一个 Excel 字段")
            read_excel_records(binding)
        except Exception as exc:
            QMessageBox.warning(self, "绑定配置不完整", str(exc))
            return
        self.accept()

    @property
    def binding(self) -> ExcelModuleBinding:
        return self._draft_binding()


class ModuleEditor(QWidget):
    changed = Signal()
    message = Signal(str)
    selection_changed = Signal(str, str)

    def __init__(self, slot_labels: dict[str, str] | None = None, parent=None):
        super().__init__(parent)
        self.slot_labels = slot_labels or {}
        self.project: PptProject | None = None
        self.manifest: TemplateManifest | None = None
        self._guard = False
        self._build_ui()

    def _button(self, text: str, slot) -> QPushButton:
        button = QPushButton(text)
        button.clicked.connect(slot)
        return button

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(7)
        hint = QLabel("模块管理页面结构；选中页面后编辑标题、小标题、图片、文字和表格。")
        hint.setObjectName("mutedLabel")
        layout.addWidget(hint)

        module_buttons = QHBoxLayout()
        module_buttons.addWidget(self._button("添加模块", self.add_module))
        module_buttons.addWidget(self._button("复制", self.duplicate_selected))
        module_buttons.addWidget(self._button("删除", self.delete_selected))
        module_buttons.addWidget(self._button("上移", lambda: self.move_selected(-1)))
        module_buttons.addWidget(self._button("下移", lambda: self.move_selected(1)))
        module_buttons.addStretch()
        module_buttons.addWidget(self._button("Excel生成", self.bind_excel))
        layout.addLayout(module_buttons)

        page_buttons = QHBoxLayout()
        page_buttons.addWidget(self._button("增加页面", self.add_page))
        page_buttons.addWidget(self._button("登记页面模板", self.register_page_template))
        page_buttons.addStretch()
        layout.addLayout(page_buttons)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        self.tree = QTreeWidget()
        self.tree.setObjectName("moduleTree")
        self.tree.setColumnCount(2)
        self.tree.setHeaderLabels(["模块 / 页面", "结构"])
        self.tree.header().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.tree.header().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        self.tree.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.tree.currentItemChanged.connect(self._show_selection)
        self.tree.itemChanged.connect(self._tree_item_changed)
        splitter.addWidget(self.tree)

        inspector = QWidget()
        inspector_layout = QVBoxLayout(inspector)
        inspector_layout.setContentsMargins(7, 0, 0, 0)
        self.empty_label = QLabel("选择一个模块或页面查看结构化内容。")
        self.empty_label.setObjectName("mutedLabel")
        inspector_layout.addWidget(self.empty_label)

        self.module_group = QGroupBox("模块属性")
        module_form = QFormLayout(self.module_group)
        self.module_name_edit = QLineEdit()
        self.module_name_edit.editingFinished.connect(self._save_module_properties)
        module_form.addRow("模块名称", self.module_name_edit)
        self.module_type_combo = QComboBox()
        for key, name in MODULE_TYPE_NAMES.items():
            self.module_type_combo.addItem(name, key)
        self.module_type_combo.currentIndexChanged.connect(self._save_module_properties)
        module_form.addRow("模块类型", self.module_type_combo)
        self.module_enabled_check = QCheckBox("参与最终 PPT 生成")
        self.module_enabled_check.stateChanged.connect(self._save_module_properties)
        module_form.addRow("启用", self.module_enabled_check)
        self.module_page_count_label = QLabel("0 页")
        module_form.addRow("当前页数", self.module_page_count_label)
        self.default_template_combo = QComboBox()
        self.default_template_combo.currentIndexChanged.connect(self._save_default_template)
        module_form.addRow("默认新增模板", self.default_template_combo)
        inspector_layout.addWidget(self.module_group)

        self.page_group = QGroupBox("页面属性")
        page_form = QFormLayout(self.page_group)
        self.page_title_edit = QLineEdit()
        self.page_title_edit.editingFinished.connect(self._save_page_properties)
        page_form.addRow("页面标题", self.page_title_edit)
        self.page_subtitle_edit = QLineEdit()
        self.page_subtitle_edit.editingFinished.connect(self._save_page_properties)
        page_form.addRow("页面小标题", self.page_subtitle_edit)
        self.page_template_combo = QComboBox()
        self.page_template_combo.currentIndexChanged.connect(self._save_page_properties)
        page_form.addRow("页面模板", self.page_template_combo)
        self.source_slide_label = QLabel("-")
        page_form.addRow("模板来源", self.source_slide_label)
        inspector_layout.addWidget(self.page_group)

        self.field_group = QGroupBox("结构化页面内容")
        field_layout = QVBoxLayout(self.field_group)
        self.field_table = QTableWidget(0, 5)
        self.field_table.setHorizontalHeaderLabels(
            ["分组", "字段", "类型", "来源", "当前值"]
        )
        self.field_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.field_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.field_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.field_table.verticalHeader().setVisible(False)
        self.field_table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.ResizeMode.ResizeToContents
        )
        self.field_table.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.ResizeMode.Stretch
        )
        self.field_table.horizontalHeader().setSectionResizeMode(
            2, QHeaderView.ResizeMode.ResizeToContents
        )
        self.field_table.horizontalHeader().setSectionResizeMode(
            3, QHeaderView.ResizeMode.ResizeToContents
        )
        self.field_table.horizontalHeader().setSectionResizeMode(
            4, QHeaderView.ResizeMode.Stretch
        )
        self.field_table.cellDoubleClicked.connect(lambda *_: self.edit_selected_field())
        field_layout.addWidget(self.field_table)
        field_buttons = QHBoxLayout()
        field_buttons.addWidget(self._button("编辑字段", self.edit_selected_field))
        field_buttons.addWidget(self._button("清除页面覆盖", self.clear_selected_override))
        field_buttons.addStretch()
        field_layout.addLayout(field_buttons)
        inspector_layout.addWidget(self.field_group, 1)
        inspector_layout.addStretch()
        splitter.addWidget(inspector)
        splitter.setSizes([280, 420])
        layout.addWidget(splitter, 1)
        self._show_empty()

    def set_state(
        self,
        project: PptProject,
        manifest: TemplateManifest,
        *,
        selection_anchor: tuple[str, int] | None = None,
    ) -> None:
        self.project = project
        self.manifest = manifest
        select_id = self._selection_id_from_anchor(selection_anchor)
        self.refresh_tree(select_id)

    def _show_empty(self) -> None:
        self.empty_label.show()
        self.module_group.hide()
        self.page_group.hide()
        self.field_group.hide()

    def _selected_ids(self) -> tuple[str, str, str]:
        item = self.tree.currentItem()
        if not item:
            return "", "", ""
        kind = str(item.data(0, Qt.ItemDataRole.UserRole) or "")
        module_id = str(item.data(0, Qt.ItemDataRole.UserRole + 1) or "")
        slide_id = str(item.data(0, Qt.ItemDataRole.UserRole + 2) or "")
        return kind, module_id, slide_id

    def selection_anchor(self) -> tuple[str, int]:
        """Return a stable module/page-position anchor for regenerated slides."""
        item = self.tree.currentItem()
        if not item:
            return "", -1
        kind, module_id, _ = self._selected_ids()
        if kind != "slide":
            return module_id, -1
        parent = item.parent()
        return module_id, parent.indexOfChild(item) if parent else 0

    def _selection_id_from_anchor(
        self, selection_anchor: tuple[str, int] | None
    ) -> str:
        if not selection_anchor or not self.project:
            return ""
        module_id, page_index = selection_anchor
        if not module_id:
            return ""
        try:
            module = module_by_id(self.project, module_id)
        except ValueError:
            return ""
        if page_index >= 0 and module.slides:
            return module.slides[min(page_index, len(module.slides) - 1)].id
        return module.id

    def _selected_module(self) -> ProjectModule | None:
        if not self.project:
            return None
        _, module_id, _ = self._selected_ids()
        if not module_id:
            return None
        try:
            return module_by_id(self.project, module_id)
        except ValueError:
            return None

    def _selected_slide(self) -> ProjectSlide | None:
        module = self._selected_module()
        if not module:
            return None
        _, _, slide_id = self._selected_ids()
        if not slide_id:
            return None
        try:
            return slide_by_id(module, slide_id)
        except ValueError:
            return None

    def refresh_tree(self, select_id: str = "") -> None:
        if not self.project:
            return
        current_kind, current_module_id, current_slide_id = self._selected_ids()
        wanted = select_id or current_slide_id or current_module_id
        self._guard = True
        try:
            self.tree.clear()
            selected_item: QTreeWidgetItem | None = None
            for module in self.project.modules:
                status = MODULE_TYPE_NAMES.get(module.module_type, module.module_type)
                if module.generated_by_binding_id:
                    status += " · Excel生成"
                elif module.module_type == "dataDriven" and not module.enabled:
                    status += " · 数据模板"
                module_item = QTreeWidgetItem(
                    [f"{module.name} · {len(module.slides)}页", status]
                )
                module_item.setData(0, Qt.ItemDataRole.UserRole, "module")
                module_item.setData(0, Qt.ItemDataRole.UserRole + 1, module.id)
                module_item.setFlags(
                    module_item.flags() | Qt.ItemFlag.ItemIsUserCheckable
                )
                module_item.setCheckState(
                    0,
                    Qt.CheckState.Checked
                    if module.enabled
                    else Qt.CheckState.Unchecked,
                )
                self.tree.addTopLevelItem(module_item)
                if module.id == wanted:
                    selected_item = module_item
                for page_index, slide in enumerate(module.slides, start=1):
                    template = page_template_by_key(module, slide.page_template_key)
                    page_item = QTreeWidgetItem(
                        [f"{page_index}. {slide.title or template.name}", f"模板第{template.source_slide}页"]
                    )
                    page_item.setData(0, Qt.ItemDataRole.UserRole, "slide")
                    page_item.setData(0, Qt.ItemDataRole.UserRole + 1, module.id)
                    page_item.setData(0, Qt.ItemDataRole.UserRole + 2, slide.id)
                    module_item.addChild(page_item)
                    if slide.id == wanted:
                        selected_item = page_item
                module_item.setExpanded(True)
            if selected_item:
                self.tree.setCurrentItem(selected_item)
            elif self.tree.topLevelItemCount():
                self.tree.setCurrentItem(self.tree.topLevelItem(0))
        finally:
            self._guard = False
        self._show_selection()

    def _tree_item_changed(self, item: QTreeWidgetItem, column: int) -> None:
        if self._guard or column != 0 or not self.project:
            return
        if item.data(0, Qt.ItemDataRole.UserRole) != "module":
            return
        module = module_by_id(
            self.project, str(item.data(0, Qt.ItemDataRole.UserRole + 1))
        )
        module.enabled = item.checkState(0) == Qt.CheckState.Checked
        sync_legacy_module_state(self.project)
        self.changed.emit()

    def _show_selection(self, *_args) -> None:
        if self._guard:
            return
        module = self._selected_module()
        if not module:
            self._show_empty()
            self.selection_changed.emit("", "")
            return
        self.empty_label.hide()
        slide = self._selected_slide()
        self._guard = True
        try:
            if slide is None:
                self.module_group.show()
                self.page_group.hide()
                self.field_group.hide()
                self.module_name_edit.setText(module.name)
                index = self.module_type_combo.findData(module.module_type)
                self.module_type_combo.setCurrentIndex(max(0, index))
                self.module_enabled_check.setChecked(module.enabled)
                self.module_page_count_label.setText(f"{len(module.slides)} 页")
                self.default_template_combo.clear()
                self.default_template_combo.addItem("添加时每次选择", "")
                for template in module.page_templates:
                    self.default_template_combo.addItem(
                        f"{template.name} · 模板第{template.source_slide}页",
                        template.key,
                    )
                default_index = self.default_template_combo.findData(
                    module.default_add_template
                )
                self.default_template_combo.setCurrentIndex(max(0, default_index))
            else:
                self.module_group.hide()
                self.page_group.show()
                self.field_group.show()
                self.page_title_edit.setText(slide.title)
                self.page_subtitle_edit.setText(slide.subtitle)
                self.page_template_combo.clear()
                for template in module.page_templates:
                    self.page_template_combo.addItem(
                        f"{template.name} · 模板第{template.source_slide}页",
                        template.key,
                    )
                template_index = self.page_template_combo.findData(
                    slide.page_template_key
                )
                self.page_template_combo.setCurrentIndex(max(0, template_index))
                template = page_template_by_key(module, slide.page_template_key)
                self.source_slide_label.setText(
                    f"模板第 {template.source_slide} 页 · {template.role}"
                )
                self._refresh_field_table(module, slide)
        finally:
            self._guard = False
        self.selection_changed.emit(module.id, slide.id if slide else "")

    def _save_module_properties(self, *_args) -> None:
        if self._guard or not self.project:
            return
        module = self._selected_module()
        if not module or self._selected_slide():
            return
        module.name = self.module_name_edit.text().strip() or module.name
        module.module_type = str(self.module_type_combo.currentData() or "fixed")
        module.enabled = self.module_enabled_check.isChecked()
        sync_legacy_module_state(self.project)
        self.refresh_tree(module.id)
        self.changed.emit()

    def _save_default_template(self, *_args) -> None:
        if self._guard:
            return
        module = self._selected_module()
        if not module or self._selected_slide():
            return
        key = str(self.default_template_combo.currentData() or "")
        if key:
            set_default_page_template(module, key)
        else:
            module.default_add_template = ""
        self.changed.emit()

    def _save_page_properties(self, *_args) -> None:
        if self._guard:
            return
        module = self._selected_module()
        slide = self._selected_slide()
        if not module or not slide:
            return
        slide.title = self.page_title_edit.text().strip()
        slide.subtitle = self.page_subtitle_edit.text().strip()
        template_key = str(self.page_template_combo.currentData() or "")
        if template_key:
            page_template_by_key(module, template_key)
            slide.page_template_key = template_key
        self.refresh_tree(slide.id)
        self.changed.emit()

    def _refresh_field_table(
        self, module: ProjectModule, slide: ProjectSlide
    ) -> None:
        if not self.project or not self.manifest:
            return
        template = page_template_by_key(module, slide.page_template_key)
        slots = slot_specs_for_source_slide(self.manifest, template.source_slide)
        self.field_table.setRowCount(len(slots))
        for row, slot in enumerate(slots):
            key = slot["key"]
            if key in slide.overrides:
                value = slide.overrides[key]
                source = "页面手工"
            elif key in module.module_values:
                value = module.module_values[key]
                source = "模块 / Excel"
            elif key in self.project.values:
                value = self.project.values[key]
                source = "项目公共"
            else:
                value = ""
                source = "模板默认"
            group = str(
                slot.get("group")
                or {"text": "文字内容", "image": "图片", "table": "表格"}.get(
                    slot["kind"], "其他"
                )
            )
            preview = self._value_preview(value, slot["kind"])
            cells = [
                QTableWidgetItem(group),
                QTableWidgetItem(self._label_for_slot(slot)),
                QTableWidgetItem(SLOT_KIND_NAMES.get(slot["kind"], slot["kind"])),
                QTableWidgetItem(source),
                QTableWidgetItem(preview),
            ]
            cells[1].setData(Qt.ItemDataRole.UserRole, key)
            for column, item in enumerate(cells):
                item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.field_table.setItem(row, column, item)

    def _value_preview(self, value: Any, kind: str) -> str:
        if kind == "table" and isinstance(value, list):
            columns = max((len(row) for row in value if isinstance(row, list)), default=0)
            return f"{len(value)}×{columns} 表格"
        if kind == "image" and value:
            return Path(str(value)).name
        text = str(value or "").replace("\n", " ")
        return text if len(text) <= 60 else text[:57] + "…"

    def _label_for_slot(self, slot: dict[str, Any]) -> str:
        return str(
            slot.get("label")
            or self.slot_labels.get(slot["key"])
            or slot["key"]
        )

    def add_module(self) -> None:
        if not self.project or not self.manifest:
            return
        choices = [
            (f"{item.get('name', item['key'])} · {len(item['slides'])}页", item["key"])
            for item in self.manifest.modules
        ]
        label, accepted = QInputDialog.getItem(
            self, "添加模块", "选择模块蓝图", [item[0] for item in choices], 0, False
        )
        if not accepted:
            return
        key = next(key for text, key in choices if text == label)
        try:
            current = self._selected_module()
            position = (
                self.project.modules.index(current) + 1 if current else len(self.project.modules)
            )
            module = add_module(
                self.project, self.manifest, key, position=position
            )
            self.refresh_tree(module.id)
            self.changed.emit()
        except Exception as exc:
            QMessageBox.warning(self, "添加模块失败", str(exc))

    def duplicate_selected(self) -> None:
        if not self.project:
            return
        module = self._selected_module()
        slide = self._selected_slide()
        if not module:
            return
        try:
            if slide:
                copied = duplicate_slide(module, slide.id)
                selected_id = copied.id
            else:
                copied_module = duplicate_module(self.project, module.id)
                selected_id = copied_module.id
            self.refresh_tree(selected_id)
            self.changed.emit()
        except Exception as exc:
            QMessageBox.warning(self, "复制失败", str(exc))

    def delete_selected(self) -> None:
        if not self.project:
            return
        module = self._selected_module()
        slide = self._selected_slide()
        if not module:
            return
        target = f"页面“{slide.title}”" if slide else f"模块“{module.name}”"
        answer = QMessageBox.question(self, "确认删除", f"确定删除{target}？\n原模板不会被修改。")
        if answer != QMessageBox.StandardButton.Yes:
            return
        try:
            if slide:
                remove_slide(module, slide.id)
                selected_id = module.id
            else:
                remove_module(self.project, module.id)
                selected_id = ""
            self.refresh_tree(selected_id)
            self.changed.emit()
        except Exception as exc:
            QMessageBox.warning(self, "删除失败", str(exc))

    def move_selected(self, offset: int) -> None:
        if not self.project:
            return
        module = self._selected_module()
        slide = self._selected_slide()
        if not module:
            return
        moved = (
            move_slide(module, slide.id, offset)
            if slide
            else move_module(self.project, module.id, offset)
        )
        if moved:
            self.refresh_tree(slide.id if slide else module.id)
            self.changed.emit()

    def add_page(self) -> None:
        module = self._selected_module()
        if not module:
            QMessageBox.information(self, "增加页面", "请先选择所属模块。")
            return
        choices = [
            (f"{item.name} · 模板第{item.source_slide}页", item.key)
            for item in module.page_templates
        ]
        if not choices:
            QMessageBox.warning(self, "增加页面", "当前模块没有可用页面模板。")
            return
        default_index = next(
            (
                index
                for index, (_, key) in enumerate(choices)
                if key == module.default_add_template
            ),
            0,
        )
        label, accepted = QInputDialog.getItem(
            self,
            "增加页面",
            "选择页面模板",
            [item[0] for item in choices],
            default_index,
            False,
        )
        if not accepted:
            return
        key = next(key for text, key in choices if text == label)
        selected_slide = self._selected_slide()
        position = (
            module.slides.index(selected_slide) + 1
            if selected_slide
            else len(module.slides)
        )
        try:
            slide = add_slide(module, key, position=position)
            self.refresh_tree(slide.id)
            self.changed.emit()
        except Exception as exc:
            QMessageBox.warning(self, "增加页面失败", str(exc))

    def register_page_template(self) -> None:
        module = self._selected_module()
        if not module or not self.manifest:
            QMessageBox.information(self, "登记页面模板", "请先选择一个模块。")
            return
        page_number, accepted = QInputDialog.getInt(
            self,
            "登记页面模板",
            "选择当前 PPT 模板中的来源页",
            1,
            1,
            self.manifest.slide_count,
        )
        if not accepted:
            return
        name, accepted = QInputDialog.getText(
            self,
            "登记页面模板",
            "页面模板名称",
            text=f"自定义模板第{page_number}页",
        )
        if not accepted:
            return
        template = add_page_template(module, page_number, name)
        set_default_page_template(module, template.key)
        self.refresh_tree(module.id)
        self.changed.emit()
        self.message.emit(f"已为模块“{module.name}”登记页面模板：{template.name}")

    def edit_selected_field(self) -> None:
        if not self.project or not self.manifest:
            return
        module = self._selected_module()
        slide = self._selected_slide()
        row = self.field_table.currentRow()
        if not module or not slide or row < 0:
            return
        key_item = self.field_table.item(row, 1)
        if not key_item:
            return
        key = str(key_item.data(Qt.ItemDataRole.UserRole) or "")
        template = page_template_by_key(module, slide.page_template_key)
        slot = next(
            (
                item
                for item in slot_specs_for_source_slide(
                    self.manifest, template.source_slide
                )
                if item["key"] == key
            ),
            None,
        )
        if not slot:
            return
        current = slide.overrides.get(
            key,
            module.module_values.get(key, self.project.values.get(key, "")),
        )
        try:
            if slot["kind"] == "text":
                dialog = TextValueDialog(
                    self._label_for_slot(slot),
                    str(current or ""),
                    slot.get("max_chars"),
                    self,
                )
                if dialog.exec() != QDialog.DialogCode.Accepted:
                    return
                value: Any = dialog.value
            elif slot["kind"] == "table":
                dialog = TableValueDialog(
                    self._label_for_slot(slot),
                    current,
                    int(slot.get("rows") or 1),
                    int(slot.get("columns") or 1),
                    self,
                )
                if dialog.exec() != QDialog.DialogCode.Accepted:
                    return
                value = dialog.value
            else:
                value = self._choose_image_value(slot)
                if value is None:
                    return
            slide.overrides[key] = value
            self._refresh_field_table(module, slide)
            self.changed.emit()
        except Exception as exc:
            QMessageBox.warning(self, "字段编辑失败", str(exc))

    def _choose_image_value(self, slot: dict[str, Any]) -> str | None:
        assert self.project is not None
        expected_category = str(slot.get("asset_category") or "")
        assets = [
            item
            for item in self.project.assets
            if not expected_category or item.category == expected_category
        ]
        choices = [f"{item.category} · {Path(item.path).name}" for item in assets]
        choices.append("浏览其他图片…")
        selected, accepted = QInputDialog.getItem(
            self, "选择图片", "项目图片素材", choices, 0, False
        )
        if not accepted:
            return None
        if selected != "浏览其他图片…":
            return assets[choices.index(selected)].path
        path, _ = QFileDialog.getOpenFileName(
            self,
            "选择图片",
            str(Path.cwd()),
            "图片 (*.png *.jpg *.jpeg *.bmp *.webp)",
        )
        return path or None

    def clear_selected_override(self) -> None:
        module = self._selected_module()
        slide = self._selected_slide()
        row = self.field_table.currentRow()
        if not module or not slide or row < 0:
            return
        item = self.field_table.item(row, 1)
        key = str(item.data(Qt.ItemDataRole.UserRole) or "") if item else ""
        if key and key in slide.overrides:
            slide.overrides.pop(key)
            self._refresh_field_table(module, slide)
            self.changed.emit()

    def bind_excel(self) -> None:
        if not self.project or not self.manifest:
            return
        module = self._selected_module()
        if not module:
            QMessageBox.information(self, "Excel 生成模块", "请先选择源模块。")
            return
        if module.generated_by_binding_id:
            QMessageBox.warning(self, "Excel 生成模块", "Excel 生成的副本不能作为数据源。")
            return
        dialog = ExcelModuleBindingDialog(
            self.project,
            module,
            self.manifest,
            self.slot_labels,
            self,
        )
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        binding = dialog.binding
        try:
            new_count = len(read_excel_records(binding))
            old_count = sum(
                item.generated_by_binding_id == binding.id
                for item in self.project.modules
            )
            pages = new_count * len(module.slides)
            answer = QMessageBox.question(
                self,
                "确认重新生成模块",
                f"旧副本：{old_count} 个\n新副本：{new_count} 个\n"
                f"当前模块结构：{len(module.slides)} 页\n预计生成：{pages} 页\n\n"
                "确认后只替换此绑定产生的旧副本。",
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
            generated = materialize_excel_modules(self.project, binding)
            self.project.excel_path = binding.source_path
            self.refresh_tree(module.id)
            self.changed.emit()
            self.message.emit(
                f"Excel 已生成 {len(generated)} 个“{module.name}”副本，共 {pages} 页。"
            )
        except Exception as exc:
            QMessageBox.critical(self, "Excel 生成模块失败", str(exc))
