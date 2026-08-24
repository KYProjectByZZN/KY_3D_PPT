from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from ppt_generator.module_service import (
    add_page_template,
    add_slide,
    duplicate_module,
    duplicate_slide,
    ensure_project_modules,
    materialize_excel_modules,
    move_slide,
    rebuild_structure_context,
    remove_module,
    remove_slide,
)
from ppt_generator.project import ExcelModuleBinding, PptProject, load_project, save_project
from ppt_generator.template_renderer import TemplateManifest


def _manifest() -> TemplateManifest:
    return TemplateManifest(
        template_filename="fixture.pptx",
        template_sha256="0" * 64,
        slide_count=4,
        modules=(
            {"key": "cover", "name": "封面", "slides": [1]},
            {"key": "optical", "name": "光学方案", "slides": [2, 3]},
            {"key": "ending", "name": "结束页", "slides": [4]},
        ),
        slots=(),
    )


class ModuleServiceTests(unittest.TestCase):
    def test_schema_v1_migrates_after_manifest_load_and_saves_as_v5(self) -> None:
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temp_dir:
            path = Path(temp_dir) / "legacy.kyppt.json"
            path.write_text(
                json.dumps(
                    {
                        "schema_version": 1,
                        "project_name": "旧项目",
                        "values": {"project_title": "保留数据"},
                        "enabled_modules": ["cover", "optical"],
                        "module_order": ["optical", "cover", "ending"],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            project = load_project(path)
            self.assertEqual(project.modules, [])
            ensure_project_modules(project, _manifest())

            self.assertEqual(
                [item.template_module_key for item in project.modules],
                ["optical", "cover", "ending"],
            )
            self.assertEqual([len(item.slides) for item in project.modules], [2, 1, 1])
            self.assertFalse(project.modules[2].enabled)
            self.assertEqual(project.values["project_title"], "保留数据")

            save_project(project, path)
            saved = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(saved["schema_version"], 5)
        self.assertEqual(len(saved["modules"]), 3)

    def test_module_and_slide_copies_are_independent(self) -> None:
        project = PptProject(values={"company": "KY"})
        manifest = _manifest()
        ensure_project_modules(project, manifest)
        source = project.modules[1]
        source.module_values["station_name"] = "工位A"
        source.slides[0].overrides["caption"] = "原图"

        page_template = add_page_template(source, 2, "补充图文页")
        added = add_slide(source, page_template.key)
        duplicate_page = duplicate_slide(source, added.id)
        self.assertEqual(len(source.slides), 4)
        self.assertNotEqual(added.id, duplicate_page.id)
        self.assertTrue(move_slide(source, duplicate_page.id, -1))
        remove_slide(source, duplicate_page.id)

        copied = duplicate_module(project, source.id)
        copied.module_values["station_name"] = "工位B"
        copied.slides[0].overrides["caption"] = "副本图"

        self.assertNotEqual(source.id, copied.id)
        self.assertEqual(len({item.id for item in copied.slides}), len(copied.slides))
        self.assertEqual(source.module_values["station_name"], "工位A")
        self.assertEqual(source.slides[0].overrides["caption"], "原图")

    def test_structure_numbers_ignore_cover_and_ending(self) -> None:
        project = PptProject()
        manifest = _manifest()
        ensure_project_modules(project, manifest)
        optical = project.modules[1]
        optical.slides[0].subtitle = "相机与光源"
        duplicate_module(project, optical.id)

        contexts = rebuild_structure_context(project, manifest)

        self.assertEqual(len(contexts), 6)
        self.assertIsNone(contexts[0].page_number)
        self.assertEqual(contexts[1].page_number, 1)
        self.assertEqual(contexts[4].page_number, 4)
        self.assertIsNone(contexts[5].page_number)
        self.assertEqual(contexts[1].total_pages, 4)
        self.assertEqual(contexts[1].module_page_count, 2)
        self.assertEqual(contexts[1].slide_subtitle, "相机与光源")

    def test_excel_rows_materialize_current_module_structure(self) -> None:
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temp_dir:
            path = Path(temp_dir) / "stations.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "光学工位"
            sheet.append(["工位名称", "相机型号"])
            sheet.append(["上料工位", "CAM-A"])
            sheet.append(["检测工位", "CAM-B"])
            sheet.append(["复检工位", "CAM-C"])
            workbook.save(path)

            project = PptProject()
            manifest = _manifest()
            ensure_project_modules(project, manifest)
            source = project.modules[1]
            custom = add_page_template(source, 2, "补充页")
            add_slide(source, custom.key)
            binding = ExcelModuleBinding(
                source_module_id=source.id,
                source_path=str(path),
                sheet="光学工位",
                data_range="A1:B4",
                field_map={"工位名称": "station_name", "相机型号": "camera_model"},
                module_name_field="工位名称",
            )

            generated = materialize_excel_modules(project, binding)
            contexts = rebuild_structure_context(project, manifest)

            self.assertFalse(source.enabled)
            self.assertEqual(len(generated), 3)
            self.assertTrue(all(len(item.slides) == 3 for item in generated))
            self.assertEqual(generated[1].module_values["camera_model"], "CAM-B")
            self.assertIn("检测工位", generated[1].name)
            self.assertEqual(len(contexts), 11)  # cover + 9 generated + ending
            self.assertEqual(len({item.slide_id for item in contexts}), len(contexts))

            remove_module(project, source.id)
            self.assertFalse(
                any(item.generated_by_binding_id == binding.id for item in project.modules)
            )
            self.assertFalse(
                any(item.id == binding.id for item in project.module_bindings)
            )


if __name__ == "__main__":
    unittest.main()
