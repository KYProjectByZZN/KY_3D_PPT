from __future__ import annotations

import base64
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QBuffer, QIODevice, Qt
from PySide6.QtGui import QColor, QPainter, QPixmap
from PySide6.QtWidgets import QTableWidgetSelectionRange
from openpyxl import Workbook

from ppt_generator import ExcelMappingRule, NavigationItem, __version__
from ppt_generator.preview import preview_fingerprint
from ppt_generator.scheme_service import materialize_equipment_scheme
from ppt_generator.ui.app import create_application
from ppt_generator.ui.dialogs import NavigationEditorDialog
from ppt_generator.ui.main_window import MainWindow


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def solid_png(width: int = 1600, height: int = 900) -> bytes:
    pixmap = QPixmap(width, height)
    pixmap.fill(QColor("#FFFFFF"))
    buffer = QBuffer()
    buffer.open(QIODevice.OpenModeFlag.WriteOnly)
    assert pixmap.save(buffer, "PNG")
    return bytes(buffer.data())


def pixmap_png(pixmap: QPixmap) -> bytes:
    buffer = QBuffer()
    buffer.open(QIODevice.OpenModeFlag.WriteOnly)
    assert pixmap.save(buffer, "PNG")
    return bytes(buffer.data())


def template_navigation_png(width: int = 1600, height: int = 900) -> bytes:
    pixmap = QPixmap(width, height)
    pixmap.fill(QColor("#FFFFFF"))
    painter = QPainter(pixmap)
    painter.fillRect(0, 106, width, 28, QColor("#D0D0D0"))
    painter.end()
    return pixmap_png(pixmap)


class DesktopUiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.application = create_application([])

    def setUp(self) -> None:
        self.window = MainWindow()
        self.window.show()
        self.application.processEvents()

    def tearDown(self) -> None:
        self.window.close()
        self.application.processEvents()

    def test_default_workspace_is_ready(self) -> None:
        self.assertEqual(
            self.window.brand_title_label.text(),
            f"AI PPT Studio v{__version__}",
        )
        self.assertEqual(self.window.tabs.count(), 6)
        self.assertEqual(self.window.tabs.tabText(5), "无CAD方案实验室")
        self.assertEqual(self.window.scheme_visual_lab_tabs.count(), 2)
        self.assertEqual(
            self.window.scheme_visual_lab_tabs.tabText(0),
            "无CAD逻辑方案",
        )
        self.assertEqual(self.window.module_workspace_tabs.count(), 2)
        self.assertEqual(self.window.module_editor.tree.topLevelItemCount(), 16)
        self.assertEqual(self.window.scheme_editor.flow_table.rowCount(), 5)
        self.assertEqual(self.window.scheme_editor.device_table.rowCount(), 0)
        self.assertEqual(
            self.window.far_generate_button.text(),
            "一键生成检测效果",
        )
        self.assertEqual(self.window.field_table.rowCount(), 67)
        self.assertEqual(self.window.structure_list.count(), 16)
        self.assertTrue(self.window.template_path_edit.text().endswith(".pptx"))
        self.assertIn("project_title", self.window.project.values)
        self.assertEqual(self.window.navigation_height_spin.value(), 0.52)
        self.assertEqual(
            self.window.navigation_font_size_combo.currentText(),
            "自动（随高度）",
        )
        self.assertEqual(
            self.window.navigation_background_label.text(),
            "当前栏目背景",
        )
        self.assertIn("#FFFFFF", self.window.navigation_color_button.text())
        self.assertEqual(
            self.window.navigation_items_label.text(),
            "公司简介 ｜ 工艺分析 ｜ 设备设计 ｜ 检测效果 ｜ 系统介绍",
        )

    def test_navigation_editor_updates_names_order_and_module_ownership(self) -> None:
        original = [
            NavigationItem("设备设计", ["equipment"]),
            NavigationItem("检测效果", ["inspection"]),
        ]
        dialog = NavigationEditorDialog(
            original,
            [
                ("equipment", "设备模块"),
                ("inspection", "检测效果"),
                ("vision", "视觉系统"),
            ],
            self.window,
        )
        dialog.add_item("系统介绍")
        dialog.set_module_keys(2, ["vision", "equipment"])
        dialog.rename_item("系统平台")
        dialog.move_item(-1)

        edited = dialog.navigation_items
        self.assertEqual(
            [item.name for item in edited],
            ["设备设计", "系统平台", "检测效果"],
        )
        self.assertEqual(edited[0].module_keys, [])
        self.assertEqual(edited[1].module_keys, ["vision", "equipment"])
        self.assertEqual(original[0].module_keys, ["equipment"])
        dialog.close()

    def test_landscape_preview_uses_top_right_adaptive_layout(self) -> None:
        self.assertEqual(
            self.window.workspace_splitter.orientation(),
            Qt.Orientation.Horizontal,
        )
        self.assertEqual(self.window.workspace_splitter.count(), 2)
        self.assertEqual(
            self.window.review_splitter.orientation(),
            Qt.Orientation.Vertical,
        )
        self.assertEqual(self.window.review_splitter.count(), 2)
        self.assertIs(
            self.window.review_splitter.widget(1),
            self.window.review_tabs,
        )
        self.assertEqual(self.window.review_tabs.count(), 2)
        self.assertEqual(self.window.review_tabs.tabText(0), "模板与输出")
        self.assertEqual(self.window.review_tabs.tabText(1), "结构与记录")
        self.assertGreaterEqual(self.window.slide_preview.canvas.minimumWidth(), 480)
        self.assertGreaterEqual(self.window.slide_preview.canvas.minimumHeight(), 270)

    def test_auto_solution_workspace_is_isolated_and_navigable(self) -> None:
        editor = self.window.auto_solution_editor
        self.assertEqual(editor.stage_list.count(), 6)
        self.assertEqual(editor.stage_stack.count(), 6)
        self.assertFalse(editor.merge_button.isEnabled())
        self.assertEqual(editor.requirement_table.columnCount(), 9)
        self.assertIsNotNone(editor.requirement_widget.current_record)
        self.assertEqual(editor.application.parser_name, "local_rule_parser_v1")
        self.assertEqual(len(self.window.project.modules), 16)

        for stage in range(6):
            editor.stage_list.setCurrentRow(stage)
            self.application.processEvents()
            self.assertEqual(editor.stage_stack.currentIndex(), stage)

        editor.stage_list.setCurrentRow(2)
        editor.add_station()
        self.application.processEvents()
        self.assertEqual(editor.station_table.rowCount(), 1)
        self.assertEqual(len(self.window.project.modules), 16)

    @patch("ppt_generator.ui.main_window.QColorDialog.getColor")
    def test_navigation_style_controls_update_project(self, get_color) -> None:
        get_color.return_value = QColor("#123456")
        cover_page = self.window.module_editor.tree.topLevelItem(0).child(0)
        self.window.module_editor.tree.setCurrentItem(cover_page)
        self.application.processEvents()
        self.window.slide_preview.set_preview(solid_png(), "测试Office", 1)

        self.window.navigation_height_spin.setValue(0.60)
        self.window.navigation_font_size_combo.setCurrentIndex(
            self.window.navigation_font_size_combo.findData(14.0)
        )
        self.window.choose_navigation_background()
        self.application.processEvents()

        self.assertEqual(
            self.window.project.presentation_style.navigation_height,
            0.60,
        )
        self.assertEqual(
            self.window.project.presentation_style.navigation_background,
            "#123456",
        )
        self.assertEqual(
            self.window.project.presentation_style.navigation_font_size,
            14.0,
        )
        self.assertIn("#123456", self.window.navigation_color_button.text())
        self.assertEqual(self.window.slide_preview._image_kind, "instant")
        image = self.window.slide_preview.canvas._source.toImage()
        scale_x = image.width() / 13.333333
        active_x = int((11.55 * scale_x / 5) * 0.08)
        self.assertEqual(image.pixelColor(active_x, 4).name(), "#123456")

    @patch("ppt_generator.ui.main_window.QMessageBox.information")
    def test_optical_far_one_click_refreshes_project_and_page_count(
        self,
        information,
    ) -> None:
        far = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "光学资料"
            / "NAT6801FAR(8.5).xlsx"
        )
        self.window.project.excel_path = str(far)

        self.window.generate_inspection_from_far()
        self.application.processEvents()

        effect_module = next(
            item
            for item in self.window.project.modules
            if item.template_module_key == "inspection_result"
        )
        self.assertEqual(len(effect_module.slides), 6)
        self.assertEqual(len(self.window.project.assets), 6)
        self.assertEqual(self.window.metric_slides.text(), "24")
        self.assertEqual(self.window.tabs.currentIndex(), 1)
        information.assert_called_once()

    def test_module_check_state_updates_project_and_preview(self) -> None:
        first = self.window.module_editor.tree.topLevelItem(0)
        module_id = first.data(0, Qt.ItemDataRole.UserRole + 1)
        first.setCheckState(0, Qt.CheckState.Unchecked)
        self.application.processEvents()

        module = next(item for item in self.window.project.modules if item.id == module_id)
        self.assertFalse(module.enabled)
        self.assertEqual(self.window.structure_list.count(), 15)
        self.assertEqual(self.window.metric_modules.text(), "15")

    def test_module_copy_and_page_structure_editor(self) -> None:
        tree = self.window.module_editor.tree
        source_item = tree.topLevelItem(5)
        source_module_id = source_item.data(0, Qt.ItemDataRole.UserRole + 1)
        tree.setCurrentItem(source_item)
        self.window.module_editor.duplicate_selected()
        self.application.processEvents()

        self.assertEqual(tree.topLevelItemCount(), 17)
        source = next(
            item for item in self.window.project.modules if item.id == source_module_id
        )
        copied = self.window.project.modules[6]
        self.assertNotEqual(source.id, copied.id)
        self.assertEqual(len(source.slides), len(copied.slides))
        self.assertTrue(
            set(item.id for item in source.slides).isdisjoint(
                item.id for item in copied.slides
            )
        )

        cover_page = tree.topLevelItem(0).child(0)
        tree.setCurrentItem(cover_page)
        self.application.processEvents()
        self.assertFalse(self.window.module_editor.page_group.isHidden())
        self.assertGreater(self.window.module_editor.field_table.rowCount(), 0)
        self.assertEqual(
            self.window.slide_preview.current_slide_id,
            self.window.project.modules[0].slides[0].id,
        )
        self.assertIn("自动化测试显示模式", self.window.slide_preview.status_label.text())
        cache_key = preview_fingerprint(
            self.window.project,
            self.window.slide_preview.current_module_id,
            self.window.slide_preview.current_slide_id,
        )
        self.window._preview_cache[cache_key] = (PNG_1X1, "测试后端", 1)
        self.window._schedule_page_preview()
        self.assertIn("内存缓存", self.window.slide_preview.status_label.text())
        self.assertIsNone(self.window._preview_worker)

    def test_excel_workbench_loads_sheet_selection_and_rule(self) -> None:
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temp_dir:
            path = Path(temp_dir) / "ui_mapping.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "项目数据"
            sheet.append(["字段", "值"])
            sheet.append(["项目标题", "界面映射测试"])
            workbook.save(path)

            self.window._load_excel_path(str(path), log=False)
            self.window.excel_preview_table.setRangeSelected(
                QTableWidgetSelectionRange(1, 1, 1, 1), True
            )
            self.window.project.excel_mappings.append(
                ExcelMappingRule("项目数据", "B2", "project_title")
            )
            self.window._refresh_excel_mapping_table()
            self.application.processEvents()

            self.assertEqual(self.window.excel_sheet_combo.currentText(), "项目数据")
            self.assertEqual(self.window._selected_excel_range(), "B2")
            self.assertEqual(self.window.excel_mapping_table.rowCount(), 1)
            self.assertIn("界面映射测试", self.window.excel_mapping_table.item(0, 5).text())

    def test_two_level_overlay_and_adjacent_preload_queue(self) -> None:
        tree = self.window.module_editor.tree
        page = tree.topLevelItem(4).child(0)
        tree.setCurrentItem(page)
        self.application.processEvents()

        pane = self.window.slide_preview
        pane.set_base_preview(PNG_1X1, 5, "设备总览")
        self.assertEqual(pane._image_kind, "base")
        self.assertTrue(pane.canvas.loading_overlay.isVisible())
        self.assertIn("模板第 5 页", pane.canvas.loading_overlay.text())
        pane.set_preview(PNG_1X1, "测试Office", 5)
        self.assertEqual(pane._image_kind, "live")
        self.assertFalse(pane.canvas.loading_overlay.isVisible())

        self.window._preload_queue.clear()
        self.window._queue_adjacent_previews(
            pane.current_module_id,
            pane.current_slide_id,
        )
        self.assertEqual(len(self.window._preload_queue), 2)
        queued_slide_ids = {item[1] for item in self.window._preload_queue}
        self.assertIn(self.window.project.modules[3].slides[0].id, queued_slide_ids)
        self.assertIn(self.window.project.modules[5].slides[0].id, queued_slide_ids)

    def test_navigation_overlay_updates_active_cell_and_height_immediately(self) -> None:
        pane = self.window.slide_preview
        pane.set_preview(solid_png(), "测试Office", 1)

        applied = pane.apply_navigation_overlay(
            ["公司简介", "工艺分析", "设备设计", "检测效果", "系统介绍"],
            0,
            0.60,
            "#123456",
            12.0,
        )

        self.assertTrue(applied)
        self.assertEqual(pane._image_kind, "instant")
        self.assertFalse(pane.canvas.loading_overlay.isVisible())
        self.assertIn("高度 0.60 in", pane.status_label.text())
        self.assertIn("字号 12 pt", pane.status_label.text())
        image = pane.canvas._source.toImage()
        scale_x = image.width() / 13.333333
        scale_y = image.height() / 7.5
        navigation_left = 0.0
        cell_width = 11.55 * scale_x / 5
        active_x = int(navigation_left + cell_width * 0.08)
        inactive_x = int(navigation_left + cell_width * 1.08)
        logo_x = int(12.8 * scale_x)
        top_y = 4
        self.assertEqual(image.pixelColor(active_x, top_y).name(), "#123456")
        self.assertEqual(image.pixelColor(inactive_x, top_y).name(), "#ffffff")

        underline_y = int(0.60 * scale_y - 1)
        self.assertEqual(image.pixelColor(active_x, underline_y).name(), "#c90000")
        self.assertEqual(image.pixelColor(inactive_x, underline_y).name(), "#d3d9de")
        self.assertEqual(image.pixelColor(logo_x, underline_y).name(), "#d3d9de")
        shadow_y = int(0.60 * scale_y + 1)
        active_shadow = image.pixelColor(active_x, shadow_y)
        inactive_shadow = image.pixelColor(inactive_x, shadow_y)
        self.assertEqual(active_shadow.name(), inactive_shadow.name())
        self.assertEqual(active_shadow.name(), "#ffffff")

        lower_y = int(0.50 * scale_y)
        self.assertEqual(image.pixelColor(active_x, lower_y).name(), "#123456")
        pane.apply_navigation_overlay(
            ["公司简介", "工艺分析", "设备设计", "检测效果", "系统介绍"],
            0,
            0.42,
            "#123456",
            9.0,
        )
        resized_image = pane.canvas._source.toImage()
        self.assertEqual(resized_image.pixelColor(active_x, lower_y).name(), "#ffffff")

    def test_height_change_repaints_selected_page_without_loading_mask(self) -> None:
        cover_page = self.window.module_editor.tree.topLevelItem(0).child(0)
        self.window.module_editor.tree.setCurrentItem(cover_page)
        self.application.processEvents()
        self.window.slide_preview.set_preview(
            template_navigation_png(),
            "测试Office",
            1,
        )
        style = self.window.project.presentation_style
        style.navigation_background = "#FFFFFF"
        style.navigation_font_size = None
        style.navigation_height = 0.52
        self.window._refresh_after_navigation_style_change()
        before = pixmap_png(self.window.slide_preview.canvas._source)
        before_image = self.window.slide_preview.canvas._source.toImage()
        scale_x = before_image.width() / 13.333333
        active_x = 8
        before_red_rows = [
            y
            for y in range(100)
            if before_image.pixelColor(active_x, y).name() == "#c90000"
        ]
        self.assertEqual(before_image.pixelColor(800, 126).name(), "#ffffff")

        self.window.navigation_height_spin.setValue(0.72)
        self.application.processEvents()

        after = pixmap_png(self.window.slide_preview.canvas._source)
        after_image = self.window.slide_preview.canvas._source.toImage()
        after_red_rows = [
            y
            for y in range(100)
            if after_image.pixelColor(active_x, y).name() == "#c90000"
        ]
        self.assertNotEqual(before, after)
        self.assertTrue(before_red_rows)
        self.assertTrue(after_red_rows)
        self.assertGreater(min(after_red_rows), min(before_red_rows) + 15)
        self.assertEqual(after_image.pixelColor(800, 126).name(), "#ffffff")
        self.assertFalse(
            self.window.slide_preview.canvas.loading_overlay.isVisible()
        )

    def test_scheme_sync_restores_regenerated_flow_page_position(self) -> None:
        flow_index, flow_module = next(
            (index, module)
            for index, module in enumerate(self.window.project.modules)
            if module.template_module_key == "inspection_flow"
        )
        flow_item = self.window.module_editor.tree.topLevelItem(flow_index)
        self.assertGreater(flow_item.childCount(), 0)
        self.window.module_editor.tree.setCurrentItem(flow_item.child(0))
        self.application.processEvents()
        old_slide_id = flow_module.slides[0].id

        self.window.project.equipment_scheme.flow_nodes[0].name = "上料定位（已修改）"
        materialize_equipment_scheme(
            self.window.project,
            self.window.manifest,
            require_module_images=False,
        )
        new_slide_id = flow_module.slides[0].id
        self.assertNotEqual(old_slide_id, new_slide_id)

        self.window._on_scheme_materialized()
        self.application.processEvents()

        self.assertEqual(
            self.window.module_editor.selection_anchor(),
            (flow_module.id, 0),
        )
        self.assertEqual(self.window.slide_preview.current_module_id, flow_module.id)
        self.assertEqual(self.window.slide_preview.current_slide_id, new_slide_id)

    def test_no_cad_scene_syncs_through_main_window_application_handler(self) -> None:
        scene = self.window.no_cad_scheme_editor.scene

        self.window._on_no_cad_scheme_committed(scene.to_dict())
        self.application.processEvents()

        imported = [
            value
            for value in self.window.project.equipment_scheme.equipment_modules
            if value.source_scene_node_id
        ]
        self.assertEqual(len(imported), len(scene.nodes))
        self.assertTrue(all(value.structure_definition for value in imported))
        self.assertTrue(all(value.image_prompt for value in imported))
        self.assertEqual(self.window.scheme_editor.device_table.columnCount(), 8)
        self.assertEqual(
            self.window.scheme_editor.device_table.horizontalHeaderItem(3).text(),
            "结构",
        )
        self.assertEqual(
            self.window.scheme_editor.device_table.item(0, 3).text(),
            "已绑定",
        )
        self.assertEqual(
            self.window.scheme_editor.device_table.item(0, 4).text(),
            "已生成",
        )
        self.assertEqual(
            self.window.scheme_editor.device_table.item(0, 5).text(),
            "待添加",
        )
        self.assertEqual(
            self.window.scheme_editor.device_table.item(0, 6).text(),
            "默认绑定",
        )
        self.assertEqual(self.window.tabs.currentIndex(), 1)
        self.assertIs(
            self.window.module_workspace_tabs.currentWidget(),
            self.window.scheme_editor,
        )

    def test_candidate_history_change_auto_saves_existing_project(self) -> None:
        target = self.window.no_cad_scheme_editor.current_result.visual_target(
            "overview"
        )
        record = {
            "schemaVersion": "project-ai-image-batch/v1",
            "projectId": self.window.project.project_id,
            "generationTarget": {
                "targetId": target.target_id,
                "targetHash": target.target_hash,
            },
            "batch": {"batchId": "batch-autosave", "candidates": []},
        }
        self.window.no_cad_scheme_editor.candidate_batch_records = [record]
        payload = self.window.no_cad_scheme_editor.workspace_snapshot()
        payload["candidateHistoryChanged"] = True
        project_file = Path("existing-project.kyppt.json")
        self.window._project_file = project_file

        with patch("ppt_generator.ui.main_window.save_project") as persist:
            persist.return_value = project_file
            self.window._on_no_cad_workspace_changed(payload)

        persist.assert_called_once_with(self.window.project, project_file)
        self.assertEqual(self.window.project.ai_image_batches, [record])
        self.assertEqual(
            self.window.project.no_cad_scene,
            self.window.no_cad_scheme_editor.scene.to_dict(),
        )

    def test_manual_device_add_creates_one_complete_binding_row(self) -> None:
        editor = self.window.scheme_editor

        editor.add_device()
        self.application.processEvents()

        self.assertEqual(editor.device_table.rowCount(), 1)
        module = self.window.project.equipment_scheme.equipment_modules[0]
        self.assertEqual(editor.device_table.item(0, 3).text(), "待定义")
        self.assertEqual(editor.device_table.item(0, 4).text(), "待生成")
        self.assertEqual(editor.device_table.item(0, 5).text(), "待添加")
        self.assertEqual(editor.device_table.item(0, 6).text(), "默认绑定")
        self.assertEqual(module.structure_definition, {})
        self.assertEqual(module.image_prompt, "")
        self.assertEqual(module.image_path, "")


if __name__ == "__main__":
    unittest.main()
